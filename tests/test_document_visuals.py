from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

from product_evidence_guard.engine import (
    _analyze_visual_with_request_cache,
    _use_pdf_observation_only_mode,
    analyze_directory,
)
from product_evidence_guard.hybrid_image_reader import (
    OBSERVATION_FACT_OVERRIDE_REASON,
    HybridImageReader,
)
from product_evidence_guard.model_output_schema import ModelOutputIssue
from product_evidence_guard.models import FactCandidate, SourceBlock
from product_evidence_guard.openvino_adapter import OpenVinoDeviceSelection
from product_evidence_guard.parsers import (
    PdfParseResult,
    RenderedPdfPage,
    RenderedPdfVisualRegion,
    parse_docx_document,
    parse_pdf_document,
    parse_xlsx_document,
    sha256_file,
)
from product_evidence_guard.qwen_vl_reader import QwenVlReadResult


_FAKE_CPU_SELECTION = OpenVinoDeviceSelection(
    requested="CPU",
    actual="CPU",
    available_devices=("CPU",),
    full_device_names={"CPU": "Mock CPU"},
    policy="explicit",
)


def _create_parameter_image(path: Path) -> None:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (900, 420), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, 860, 380), outline="black", width=4)
    draw.text((80, 100), "INPUT VOLTAGE: 24V", fill="black")
    draw.text((80, 180), "POWER: 45W", fill="black")
    draw.line((100, 310, 800, 310), fill="black", width=3)
    draw.text((380, 325), "120 mm", fill="black")
    image.save(path, format="PNG")


def _create_docx_with_image(path: Path, image_path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_paragraph("Embedded engineering label")
    run = document.add_paragraph().add_run()
    run.add_picture(str(image_path), width=Inches(4.5))
    document.save(path)


def _create_xlsx_with_image_and_chart(path: Path, image_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as SpreadsheetImage

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Performance"
    sheet.append(["Mode", "Power (W)"])
    sheet.append(["Eco", 30])
    sheet.append(["Turbo", 65])

    chart = BarChart()
    chart.title = "Power by mode"
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=3),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(sheet, min_col=1, min_row=2, max_row=3)
    )
    sheet.add_chart(chart, "D2")
    sheet.add_image(SpreadsheetImage(str(image_path)), "D18")
    workbook.save(path)
    workbook.close()


def _create_xlsx_with_image(path: Path, image_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SpreadsheetImage

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Visual"
    sheet.add_image(SpreadsheetImage(str(image_path)), "C4")
    workbook.save(path)
    workbook.close()


class _LargeImageObject(dict):
    def get_object(self):
        return self


class _MixedPdfPage:
    def extract_text(self) -> str:
        return "Mechanical drawing\nAll dimensions in mm"

    def get(self, key: str):
        if key != "/Resources":
            return None
        return _LargeImageObject(
            {
                "/XObject": _LargeImageObject(
                    {
                        "/Im1": _LargeImageObject(
                            {
                                "/Subtype": "/Image",
                                "/Width": 1000,
                                "/Height": 600,
                            }
                        )
                    }
                )
            }
        )

    def get_contents(self):
        return types.SimpleNamespace(operations=[])


class _VectorPdfPage:
    def __init__(self, text: str, operator_count: int = 500) -> None:
        self._text = text
        self._operator_count = operator_count

    def extract_text(self) -> str:
        return self._text

    def get(self, key: str):
        if key == "/Resources":
            return _LargeImageObject({"/XObject": _LargeImageObject({})})
        return None

    def get_contents(self):
        return types.SimpleNamespace(
            operations=[([], b"l")] * self._operator_count
        )


class _NestedImagePdfPage(_VectorPdfPage):
    def __init__(self) -> None:
        super().__init__("Application curves\nFigure 18", operator_count=2)

    def get(self, key: str):
        if key != "/Resources":
            return None
        nested_image = _LargeImageObject(
            {
                "/Subtype": "/Image",
                "/Width": 1023,
                "/Height": 709,
            }
        )
        form = _LargeImageObject(
            {
                "/Subtype": "/Form",
                "/Resources": _LargeImageObject(
                    {
                        "/XObject": _LargeImageObject(
                            {"/Scope": nested_image}
                        )
                    }
                ),
            }
        )
        return _LargeImageObject(
            {"/XObject": _LargeImageObject({"/ChartForm": form})}
        )


class _CountingVisualReader:
    def __init__(self) -> None:
        self.calls = 0
        self.seen_paths: list[Path] = []

    def analyze_image(
        self,
        image_path: Path,
        root: Path,
        *,
        deadline: float | None = None,
    ) -> QwenVlReadResult:
        self.calls += 1
        self.seen_paths.append(Path(image_path))
        relative = Path(image_path).relative_to(root).as_posix()
        image_hash = sha256_file(Path(image_path))
        block = SourceBlock(
            block_id=f"temporary-block-{self.calls}",
            source_file=relative,
            source_kind="image_openvino_rapidocr",
            file_hash=image_hash,
            locator={
                "image": relative,
                "transcription_id": "ocr-001",
                "bbox_1000": [80, 200, 700, 300],
                "position_precision": "approximate",
                "recognition_confidence_source": (
                    "rapidocr_recognizer_score"
                ),
            },
            text="INPUT VOLTAGE: 24V",
            recognition_confidence=0.99,
            extraction_method="rapidocr_openvino_transcription",
            recognition_confidence_source="rapidocr_recognizer_score",
        )
        candidate = FactCandidate(
            candidate_id=f"temporary-candidate-{self.calls}",
            field="voltage",
            field_label="电压",
            raw_value="24V",
            normalized_value=24,
            normalized_unit="V",
            source_block_id=block.block_id,
            source_file=relative,
            source_kind=block.source_kind,
            file_hash=image_hash,
            locator=dict(block.locator),
            raw_text=block.text,
            recognition_confidence=0.99,
            mapping_confidence=0.99,
            extraction_method="rapidocr_deterministic_mapping",
            scope="input",
            mapping_confidence_source="deterministic",
        )
        return QwenVlReadResult(
            image_file=relative,
            file_hash=image_hash,
            source_blocks=[block],
            fact_candidates=[candidate],
            image_route="ocr_fast",
            ocr_line_count=1,
        )


class _ObservationOverrideReader(HybridImageReader):
    def __init__(self) -> None:
        self.calls = 0

    def analyze_image(
        self,
        image_path: Path,
        root: Path,
        *,
        deadline: float | None = None,
        observation_only: bool = False,
    ) -> QwenVlReadResult:
        del deadline
        self.calls += 1
        if not observation_only:
            raise AssertionError("test must select observation mode first")
        relative = Path(image_path).relative_to(root).as_posix()
        image_hash = sha256_file(Path(image_path))
        block = SourceBlock(
            block_id="override-image-block",
            source_file=relative,
            source_kind="image_openvino_rapidocr",
            file_hash=image_hash,
            locator={
                "image": relative,
                "transcription_id": "ocr-001",
                "bbox_1000": [80, 200, 700, 300],
                "position_precision": "approximate",
            },
            text="NET WEIGHT: 320 g",
            recognition_confidence=0.99,
            extraction_method="rapidocr_openvino_transcription",
            recognition_confidence_source="rapidocr_recognizer_score",
        )
        candidate = FactCandidate(
            candidate_id="override-image-candidate",
            field="net_weight",
            field_label="净重",
            raw_value="320 g",
            normalized_value=320,
            normalized_unit="g",
            source_block_id=block.block_id,
            source_file=relative,
            source_kind=block.source_kind,
            file_hash=image_hash,
            locator=dict(block.locator),
            raw_text=block.text,
            recognition_confidence=0.99,
            mapping_confidence=0.99,
            extraction_method="rapidocr_deterministic_mapping",
            scope="net",
            mapping_confidence_source="deterministic",
        )
        return QwenVlReadResult(
            image_file=relative,
            file_hash=image_hash,
            source_blocks=[block],
            fact_candidates=[candidate],
            image_route="ocr_fast",
            fallback_reasons=[OBSERVATION_FACT_OVERRIDE_REASON],
            ocr_line_count=1,
        )


class _ErrorVisualReader:
    def __init__(self, *, code: str = "backend_generation_failed") -> None:
        self.calls = 0
        self.code = code

    def analyze_image(
        self,
        image_path: Path,
        root: Path,
        *,
        deadline: float | None = None,
    ) -> QwenVlReadResult:
        del deadline
        self.calls += 1
        relative = Path(image_path).relative_to(root).as_posix()
        return QwenVlReadResult(
            image_file=relative,
            file_hash=sha256_file(Path(image_path)),
            errors=[
                ModelOutputIssue(
                    stage="visual_transcription",
                    code=self.code,
                    message="controlled failure",
                )
            ],
            image_route="qwen_deep_fallback",
        )


class _RaisingVisualReader:
    def __init__(self) -> None:
        self.calls = 0

    def analyze_image(
        self,
        image_path: Path,
        root: Path,
        *,
        deadline: float | None = None,
    ) -> QwenVlReadResult:
        del image_path, root, deadline
        self.calls += 1
        raise RuntimeError("controlled reader exception")


class DocumentVisualParserTests(unittest.TestCase):
    def test_docx_embedded_image_is_extracted_with_paragraph_locator(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            image = root / "drawing.png"
            document = root / "drawing.docx"
            _create_parameter_image(image)
            _create_docx_with_image(document, image)

            result = parse_docx_document(
                document,
                document.name,
                sha256_file(document),
            )

            self.assertEqual(len(result.visual_assets), 1)
            asset = result.visual_assets[0]
            self.assertEqual(asset.source_kind, "docx_embedded_image")
            self.assertEqual(asset.extension, ".png")
            self.assertEqual(asset.locator["paragraph"], 1)
            self.assertEqual(
                asset.content_hash,
                hashlib.sha256(asset.data).hexdigest(),
            )
            self.assertFalse(result.visual_issues)

    def test_xlsx_image_and_native_chart_are_both_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            image = root / "drawing.png"
            workbook = root / "performance.xlsx"
            _create_parameter_image(image)
            _create_xlsx_with_image_and_chart(workbook, image)

            result = parse_xlsx_document(
                workbook,
                workbook.name,
                sha256_file(workbook),
            )

            self.assertEqual(len(result.visual_assets), 1)
            self.assertEqual(
                result.visual_assets[0].locator["sheet"],
                "Performance",
            )
            self.assertEqual(len(result.structured_visuals), 1)
            chart = result.structured_visuals[0]
            self.assertEqual(chart.title, "Power by mode")
            self.assertEqual(chart.status, "structured")
            series = chart.structured_data["series"][0]
            self.assertEqual(series["name"], "Power (W)")
            self.assertEqual(
                series["points"],
                [
                    {"category": "Eco", "value": 30},
                    {"category": "Turbo", "value": 65},
                ],
            )

    def test_pdf_page_with_text_and_large_image_is_marked_mixed(self) -> None:
        fake_reader = types.SimpleNamespace(pages=[_MixedPdfPage()])
        fake_pypdf = types.SimpleNamespace(PdfReader=lambda path: fake_reader)
        with patch.dict(sys.modules, {"pypdf": fake_pypdf}):
            result = parse_pdf_document(
                Path("drawing.pdf"),
                "drawing.pdf",
                "pdf-hash",
            )

        self.assertEqual(result.scanned_page_numbers, ())
        self.assertEqual(result.mixed_visual_page_numbers, (1,))
        self.assertEqual(
            result.visual_page_reasons[0]["reason"],
            "large_embedded_image",
        )

    def test_dense_decorative_vector_page_is_not_sent_to_image_model(self) -> None:
        fake_reader = types.SimpleNamespace(
            pages=[
                _VectorPdfPage(
                    "Raspberry Pi 45W USB-C Power Supply\nPublished April 2025"
                )
            ]
        )
        fake_pypdf = types.SimpleNamespace(PdfReader=lambda path: fake_reader)
        with patch.dict(sys.modules, {"pypdf": fake_pypdf}):
            result = parse_pdf_document(
                Path("cover.pdf"),
                "cover.pdf",
                "pdf-hash",
            )

        self.assertEqual(result.mixed_visual_page_numbers, ())

    def test_vector_physical_specification_page_is_marked_mixed(self) -> None:
        fake_reader = types.SimpleNamespace(
            pages=[
                _VectorPdfPage(
                    "Physical specification Type G\nAll dimensions in mm"
                )
            ]
        )
        fake_pypdf = types.SimpleNamespace(PdfReader=lambda path: fake_reader)
        with patch.dict(sys.modules, {"pypdf": fake_pypdf}):
            result = parse_pdf_document(
                Path("drawing.pdf"),
                "drawing.pdf",
                "pdf-hash",
            )

        self.assertEqual(result.mixed_visual_page_numbers, (1,))
        self.assertEqual(
            result.visual_page_reasons[0]["reason"],
            "keyword_with_vector_drawing",
        )

    def test_image_nested_in_pdf_form_is_marked_mixed(self) -> None:
        fake_reader = types.SimpleNamespace(pages=[_NestedImagePdfPage()])
        fake_pypdf = types.SimpleNamespace(PdfReader=lambda path: fake_reader)
        with patch.dict(sys.modules, {"pypdf": fake_pypdf}):
            result = parse_pdf_document(
                Path("nested-image.pdf"),
                "nested-image.pdf",
                "pdf-hash",
            )

        self.assertEqual(result.mixed_visual_page_numbers, (1,))
        self.assertEqual(
            result.visual_page_reasons[0]["reason"],
            "large_embedded_image",
        )
        self.assertEqual(
            result.visual_page_reasons[0]["large_image_count"],
            1,
        )


class DocumentVisualEngineTests(unittest.TestCase):
    def test_identical_docx_and_xlsx_images_share_request_cache(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            image = root / "shared.png"
            document = input_dir / "a-shared.docx"
            workbook = input_dir / "b-shared.xlsx"
            _create_parameter_image(image)
            _create_docx_with_image(document, image)
            _create_xlsx_with_image(workbook, image)
            reader = _CountingVisualReader()

            summary = analyze_directory(
                input_dir,
                output_dir,
                openvino_vlm_model=str(root / "fake-model"),
                preloaded_image_reader=reader,
                preloaded_model_load_seconds=0.0,
                preloaded_model_reused=True,
                device_selection=_FAKE_CPU_SELECTION,
            )

            self.assertEqual(reader.calls, 1)
            self.assertFalse(reader.seen_paths[0].exists())
            self.assertEqual(
                summary["visual_content_cache"],
                {
                    "hits": 1,
                    "misses": 1,
                    "avoided_reader_calls": 1,
                },
            )
            facts = json.loads(
                (output_dir / "product-facts.json").read_text(
                    encoding="utf-8"
                )
            )
            candidates = {
                item["source_file"]: item
                for item in facts["candidates"]
                if "embedded_image" in item["source_kind"]
            }
            self.assertEqual(
                set(candidates),
                {"a-shared.docx", "b-shared.xlsx"},
            )
            docx_candidate = candidates["a-shared.docx"]
            xlsx_candidate = candidates["b-shared.xlsx"]
            self.assertNotEqual(
                docx_candidate["candidate_id"],
                xlsx_candidate["candidate_id"],
            )
            self.assertNotEqual(
                docx_candidate["source_block_id"],
                xlsx_candidate["source_block_id"],
            )
            self.assertEqual(docx_candidate["locator"]["paragraph"], 1)
            self.assertEqual(xlsx_candidate["locator"]["sheet"], "Visual")
            self.assertEqual(
                xlsx_candidate["locator"]["anchor"],
                {"row": 4, "column": 3},
            )
            self.assertEqual(
                docx_candidate["file_hash"],
                sha256_file(document),
            )
            self.assertEqual(
                xlsx_candidate["file_hash"],
                sha256_file(workbook),
            )

            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )
            embedded = {
                item["source_file"]: item
                for item in visuals["visuals"]
                if item["asset_type"] == "embedded_image"
            }
            self.assertFalse(embedded["a-shared.docx"]["visual_cache_hit"])
            self.assertTrue(embedded["b-shared.xlsx"]["visual_cache_hit"])
            self.assertEqual(
                embedded["a-shared.docx"]["visual_content_hash"],
                embedded["b-shared.xlsx"]["visual_content_hash"],
            )
            self.assertEqual(
                embedded["a-shared.docx"]["visual_cache_scope"],
                "request",
            )
            transcription = json.loads(
                (output_dir / "visual-transcription.json").read_text(
                    encoding="utf-8"
                )
            )
            visual_results = {
                item["source_file"]: item
                for item in transcription["images"]
            }
            self.assertFalse(
                visual_results["a-shared.docx"]["visual_cache_hit"]
            )
            self.assertTrue(
                visual_results["b-shared.xlsx"]["visual_cache_hit"]
            )
            self.assertEqual(
                visual_results["a-shared.docx"]["visual_content_hash"],
                visual_results["b-shared.xlsx"]["visual_content_hash"],
            )

    def test_visual_errors_are_not_cached_and_retry_next_run(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            image = root / "shared.png"
            _create_parameter_image(image)
            _create_docx_with_image(input_dir / "a.docx", image)
            _create_xlsx_with_image(input_dir / "b.xlsx", image)
            reader = _ErrorVisualReader()
            analyze_options = {
                "openvino_vlm_model": str(root / "fake-model"),
                "preloaded_image_reader": reader,
                "preloaded_model_load_seconds": 0.0,
                "preloaded_model_reused": True,
                "device_selection": _FAKE_CPU_SELECTION,
            }

            first = analyze_directory(
                input_dir,
                output_dir,
                **analyze_options,
            )
            second = analyze_directory(
                input_dir,
                output_dir,
                **analyze_options,
            )

            self.assertEqual(reader.calls, 4)
            self.assertEqual(
                first["visual_content_cache"],
                {
                    "hits": 0,
                    "misses": 2,
                    "avoided_reader_calls": 0,
                },
            )
            self.assertEqual(
                second["visual_content_cache"],
                {
                    "hits": 0,
                    "misses": 2,
                    "avoided_reader_calls": 0,
                },
            )
            self.assertFalse(second["unchanged_files_reused"])
            self.assertEqual(
                second["changed_or_new_files"],
                ["a.docx", "b.xlsx"],
            )
            state = json.loads(
                (output_dir / "analysis-state.json").read_text(
                    encoding="utf-8"
                )
            )
            self.assertTrue(state["files"]["a.docx"]["retry_required"])
            self.assertTrue(state["files"]["b.xlsx"]["retry_required"])
            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )
            embedded = [
                item
                for item in visuals["visuals"]
                if item["asset_type"] == "embedded_image"
            ]
            self.assertEqual(len(embedded), 2)
            self.assertTrue(
                all(
                    item["status"] == "processed_with_errors"
                    and item["visual_cache_hit"] is False
                    for item in embedded
                )
            )

    def test_timeouts_and_exceptions_do_not_populate_visual_cache(self) -> None:
        readers = (
            _ErrorVisualReader(code="file_timeout"),
            _RaisingVisualReader(),
        )
        for reader in readers:
            with self.subTest(reader=type(reader).__name__):
                with tempfile.TemporaryDirectory() as temporary:
                    root = Path(temporary)
                    input_dir = root / "input"
                    output_dir = root / "output"
                    input_dir.mkdir()
                    image = root / "shared.png"
                    _create_parameter_image(image)
                    _create_docx_with_image(input_dir / "a.docx", image)
                    _create_xlsx_with_image(input_dir / "b.xlsx", image)

                    summary = analyze_directory(
                        input_dir,
                        output_dir,
                        openvino_vlm_model=str(root / "fake-model"),
                        preloaded_image_reader=reader,
                        preloaded_model_load_seconds=0.0,
                        preloaded_model_reused=True,
                        device_selection=_FAKE_CPU_SELECTION,
                    )

                    self.assertEqual(reader.calls, 2)
                    self.assertEqual(
                        summary["visual_content_cache"],
                        {
                            "hits": 0,
                            "misses": 2,
                            "avoided_reader_calls": 0,
                        },
                    )
                    state = json.loads(
                        (output_dir / "analysis-state.json").read_text(
                            encoding="utf-8"
                        )
                    )
                    self.assertTrue(
                        state["files"]["a.docx"]["retry_required"]
                    )
                    self.assertTrue(
                        state["files"]["b.xlsx"]["retry_required"]
                    )

    def test_observation_cache_boundary_never_calls_plain_reader(self) -> None:
        reader = _CountingVisualReader()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            image = root / "chart.png"
            _create_parameter_image(image)
            cache: dict[str, QwenVlReadResult] = {}
            stats = {"hits": 0, "misses": 0, "avoided_reader_calls": 0}

            result, cache_hit, _ = _analyze_visual_with_request_cache(
                reader,
                image,
                root,
                deadline=None,
                cache_namespace="test",
                cache=cache,
                stats=stats,
                observation_only=True,
            )

        self.assertEqual(reader.calls, 0)
        self.assertFalse(cache_hit)
        self.assertEqual(result.image_route, "observations_unavailable")
        self.assertEqual(result.fact_candidates, [])
        self.assertEqual(result.mappings, [])
        self.assertEqual(result.errors[0].code, "observations_unavailable")
        self.assertEqual(cache, {})

    def test_visual_result_returned_after_deadline_is_discarded(self) -> None:
        reader = _CountingVisualReader()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            image = root / "late.png"
            _create_parameter_image(image)
            cache: dict[str, QwenVlReadResult] = {}
            stats = {"hits": 0, "misses": 0, "avoided_reader_calls": 0}

            with patch(
                "product_evidence_guard.engine.time.perf_counter",
                side_effect=[1.0, 2.0],
            ):
                result, cache_hit, _ = _analyze_visual_with_request_cache(
                    reader,
                    image,
                    root,
                    deadline=1.5,
                    cache_namespace="test",
                    cache=cache,
                    stats=stats,
                )

        self.assertEqual(reader.calls, 1)
        self.assertFalse(cache_hit)
        self.assertEqual(result.image_route, "deadline_exceeded")
        self.assertEqual(result.fact_candidates, [])
        self.assertEqual(result.source_blocks, [])
        self.assertEqual(result.errors[0].code, "deadline_exceeded")
        self.assertEqual(cache, {})

    def test_observation_mode_requires_chart_semantics_not_just_text_size(
        self,
    ) -> None:
        rich_context = {"line_count": 32, "text_char_count": 593}
        self.assertTrue(
            _use_pdf_observation_only_mode(
                page_mode="mixed",
                reason={"keyword": "figure"},
                text_layer_context=rich_context,
            )
        )
        self.assertTrue(
            _use_pdf_observation_only_mode(
                page_mode="mixed",
                reason={"keyword": "curve"},
                text_layer_context=rich_context,
            )
        )
        self.assertFalse(
            _use_pdf_observation_only_mode(
                page_mode="mixed",
                reason={"keyword": "diagram"},
                text_layer_context=rich_context,
            )
        )
        self.assertFalse(
            _use_pdf_observation_only_mode(
                page_mode="mixed",
                reason={"keyword": "figure"},
                text_layer_context={
                    "line_count": 2,
                    "text_char_count": 40,
                },
            )
        )

    def test_embedded_docx_image_is_retargeted_and_then_cached(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            image = root / "drawing.png"
            document = input_dir / "drawing.docx"
            _create_parameter_image(image)
            _create_docx_with_image(document, image)
            reader = _CountingVisualReader()
            model_path = root / "fake-model"

            first = analyze_directory(
                input_dir,
                output_dir,
                openvino_vlm_model=str(model_path),
                preloaded_image_reader=reader,
                preloaded_model_load_seconds=0.0,
                preloaded_model_reused=True,
                device_selection=_FAKE_CPU_SELECTION,
            )
            second = analyze_directory(
                input_dir,
                output_dir,
                openvino_vlm_model=str(model_path),
                preloaded_image_reader=reader,
                preloaded_model_load_seconds=0.0,
                preloaded_model_reused=True,
                device_selection=_FAKE_CPU_SELECTION,
            )

            self.assertEqual(reader.calls, 1)
            self.assertFalse(reader.seen_paths[0].exists())
            self.assertEqual(first["candidate_count"], 1)
            self.assertEqual(
                first["document_visuals"]["embedded_images_processed"],
                1,
            )
            self.assertEqual(
                second["unchanged_files_reused"],
                ["drawing.docx"],
            )
            facts = json.loads(
                (output_dir / "product-facts.json").read_text(
                    encoding="utf-8"
                )
            )
            candidate = facts["candidates"][0]
            self.assertEqual(candidate["source_file"], "drawing.docx")
            self.assertEqual(
                candidate["source_kind"],
                "docx_embedded_image_openvino_rapidocr",
            )
            self.assertEqual(candidate["locator"]["paragraph"], 1)
            self.assertIn("asset_id", candidate["locator"])
            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )
            self.assertEqual(
                visuals["status"],
                "observations_not_facts",
            )
            self.assertEqual(
                visuals["visuals"][0]["status"],
                "processed",
            )
            markdown = (output_dir / "conflicts.md").read_text(
                encoding="utf-8"
            )
            self.assertIn("文档视觉内容", markdown)

    def test_xlsx_chart_is_structured_without_loading_image_model(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            image = root / "drawing.png"
            workbook = input_dir / "performance.xlsx"
            _create_parameter_image(image)
            _create_xlsx_with_image_and_chart(workbook, image)

            summary = analyze_directory(input_dir, output_dir)
            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )

            self.assertEqual(
                summary["document_visuals"]["native_charts_structured"],
                1,
            )
            self.assertEqual(
                summary["document_visuals"]["embedded_images_skipped"],
                1,
            )
            chart = next(
                item
                for item in visuals["visuals"]
                if item["asset_type"] == "native_chart"
            )
            self.assertEqual(
                chart["structured_data"]["series"][0]["points"][1],
                {"category": "Turbo", "value": 65},
            )

    def test_vector_pdf_uses_text_layer_without_image_model_call(self) -> None:
        block = SourceBlock(
            block_id="pdf-text-1",
            source_file="drawing.pdf",
            source_kind="pdf_text",
            file_hash="pdf-hash",
            locator={"page": 1, "text_block": 0},
            text="Physical specification | All dimensions in mm | 85 | 58",
        )
        pdf_result = PdfParseResult(
            blocks=(block,),
            scanned_page_numbers=(),
            page_count=1,
            mixed_visual_page_numbers=(1,),
            visual_page_reasons=(
                {
                    "page": 1,
                    "page_mode": "mixed",
                    "reason": "keyword_with_vector_drawing",
                    "large_image_count": 0,
                    "vector_operator_count": 5800,
                    "keyword": "dimension",
                },
            ),
        )
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "drawing.pdf").write_bytes(b"fake pdf")
            with (
                patch(
                    "product_evidence_guard.engine.parse_pdf_document",
                    return_value=pdf_result,
                ),
                patch(
                    "product_evidence_guard.engine.render_pdf_page",
                    side_effect=AssertionError("must not render vector page"),
                ),
            ):
                summary = analyze_directory(input_dir, output_dir)

            self.assertEqual(
                summary["mixed_pdf"]["pages_structured_from_text_layer"],
                1,
            )
            self.assertEqual(
                summary["mixed_pdf"]["pages_processed_with_local_image_ai"],
                0,
            )
            self.assertEqual(summary["mixed_pdf"]["pages_skipped"], 0)
            self.assertFalse(summary["skipped_files"])
            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )
            record = visuals["visuals"][0]
            self.assertEqual(
                record["status"],
                "structured_from_text_layer",
            )
            self.assertTrue(record["model_call_avoided"])
            self.assertEqual(
                record["text_layer_context"]["lines"],
                [block.text],
            )

    def test_mixed_pdf_page_keeps_original_page_evidence(self) -> None:
        reader = _CountingVisualReader()
        rendered_paths: list[Path] = []
        pdf_result = PdfParseResult(
            blocks=(),
            scanned_page_numbers=(),
            page_count=1,
            mixed_visual_page_numbers=(1,),
            visual_page_reasons=(
                {
                    "page": 1,
                    "page_mode": "mixed",
                    "reason": "large_embedded_image",
                    "large_image_count": 1,
                },
            ),
        )

        def fake_render(
            path: Path,
            page_number: int,
            output_dir: Path,
        ) -> RenderedPdfPage:
            image_path = output_dir / "page-0001.png"
            _create_parameter_image(image_path)
            rendered_paths.append(image_path)
            rendered_hashes.append(sha256_file(image_path))
            return RenderedPdfVisualRegion(
                page_number=1,
                image_path=image_path,
                pixel_width=800,
                pixel_height=600,
                crop_box_pdf=(100.0, 100.0, 500.0, 400.0),
                crop_ratio=0.25,
                strategy="pdfium_top_level_raster_union_v1",
                fallback_reason=None,
                page_width_pdf=600.0,
                page_height_pdf=800.0,
            )

        rendered_hashes: list[str] = []
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            pdf = input_dir / "mixed.pdf"
            pdf.write_bytes(b"fake pdf")
            with (
                patch(
                    "product_evidence_guard.engine.parse_pdf_document",
                    return_value=pdf_result,
                ),
                patch(
                    "product_evidence_guard.engine.render_pdf_visual_region",
                    side_effect=fake_render,
                ),
            ):
                summary = analyze_directory(
                    input_dir,
                    output_dir,
                    openvino_vlm_model=str(root / "fake-model"),
                    preloaded_image_reader=reader,
                    preloaded_model_load_seconds=0.0,
                    preloaded_model_reused=True,
                    device_selection=_FAKE_CPU_SELECTION,
                )

            facts = json.loads(
                (output_dir / "product-facts.json").read_text(
                    encoding="utf-8"
                )
            )
            candidate = facts["candidates"][0]
            self.assertEqual(
                candidate["source_kind"],
                "pdf_mixed_page_openvino_rapidocr",
            )
            self.assertEqual(candidate["locator"]["page"], 1)
            self.assertEqual(candidate["locator"]["page_mode"], "mixed")
            self.assertEqual(
                candidate["locator"]["page_source"],
                "rendered_mixed_page_roi",
            )
            self.assertEqual(
                candidate["locator"]["crop_bbox_1000"],
                [80, 200, 700, 300],
            )
            self.assertEqual(
                candidate["locator"]["bbox_1000"],
                [220, 575, 633, 612],
            )
            self.assertEqual(
                candidate["locator"]["crop_strategy"],
                "pdfium_top_level_raster_union_v1",
            )
            self.assertEqual(candidate["locator"]["crop_ratio"], 0.25)
            self.assertEqual(
                summary["mixed_pdf"]["pages_processed_with_local_image_ai"],
                1,
            )
            self.assertEqual(
                summary["visual_content_cache"],
                {
                    "hits": 0,
                    "misses": 1,
                    "avoided_reader_calls": 0,
                },
            )
            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )
            page_visual = visuals["visuals"][0]
            self.assertFalse(page_visual["visual_cache_hit"])
            self.assertEqual(
                page_visual["visual_cache_scope"],
                "request",
            )
            self.assertEqual(
                page_visual["visual_content_hash"],
                rendered_hashes[0],
            )
            self.assertEqual(
                page_visual["raw_coordinate_space"],
                "rendered_crop_image_1000",
            )
            self.assertEqual(
                page_visual["transcription_coordinate_space"],
                "pdf_page_1000",
            )
            self.assertEqual(
                page_visual["ocr_observation_coordinate_space"],
                "pdf_page_1000",
            )
        self.assertFalse(rendered_paths[0].exists())

    def test_rich_chart_page_explicit_image_fact_still_blocks_conflict(
        self,
    ) -> None:
        reader = _ObservationOverrideReader()
        blocks = [
            SourceBlock(
                block_id="pdf-weight",
                source_file="mixed.pdf",
                source_kind="pdf_text",
                file_hash="pdf-hash",
                locator={"page": 1, "text_block": 0},
                text="NET WEIGHT: 350 g",
            )
        ]
        blocks.extend(
            SourceBlock(
                block_id=f"pdf-context-{index}",
                source_file="mixed.pdf",
                source_kind="pdf_text",
                file_hash="pdf-hash",
                locator={"page": 1, "text_block": index},
                text=(
                    f"Figure context line {index}: efficiency curve "
                    "test conditions and axis notes."
                ),
            )
            for index in range(1, 13)
        )
        pdf_result = PdfParseResult(
            blocks=tuple(blocks),
            scanned_page_numbers=(),
            page_count=1,
            mixed_visual_page_numbers=(1,),
            visual_page_reasons=(
                {
                    "page": 1,
                    "page_mode": "mixed",
                    "reason": "large_embedded_image",
                    "large_image_count": 1,
                    "keyword": "figure",
                },
            ),
        )

        def fake_render(
            path: Path,
            page_number: int,
            output_dir: Path,
        ) -> RenderedPdfVisualRegion:
            del path, page_number
            image_path = output_dir / "page-0001-roi.png"
            _create_parameter_image(image_path)
            return RenderedPdfVisualRegion(
                page_number=1,
                image_path=image_path,
                pixel_width=800,
                pixel_height=600,
                crop_box_pdf=(100.0, 100.0, 500.0, 400.0),
                crop_ratio=0.25,
                strategy="pdfium_top_level_raster_union_v1",
                fallback_reason=None,
                page_width_pdf=600.0,
                page_height_pdf=800.0,
            )

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "mixed.pdf").write_bytes(b"fake pdf")
            with (
                patch(
                    "product_evidence_guard.engine.parse_pdf_document",
                    return_value=pdf_result,
                ),
                patch(
                    "product_evidence_guard.engine.render_pdf_visual_region",
                    side_effect=fake_render,
                ),
            ):
                summary = analyze_directory(
                    input_dir,
                    output_dir,
                    openvino_vlm_model=str(root / "fake-model"),
                    preloaded_image_reader=reader,
                    preloaded_model_load_seconds=0.0,
                    preloaded_model_reused=True,
                    device_selection=_FAKE_CPU_SELECTION,
                )

            facts = json.loads(
                (output_dir / "product-facts.json").read_text(
                    encoding="utf-8"
                )
            )
            visuals = json.loads(
                (output_dir / "document-visuals.json").read_text(
                    encoding="utf-8"
                )
            )

        self.assertEqual(reader.calls, 1)
        self.assertEqual(summary["blocking_conflict_count"], 1)
        weight_group = next(
            item
            for item in facts["fact_groups"]
            if item["field"] == "net_weight"
        )
        self.assertEqual(weight_group["classification"], "strong_conflict")
        self.assertEqual(weight_group["severity"], "block")
        self.assertEqual(
            {item["normalized_value"] for item in facts["candidates"]},
            {320, 350},
        )
        page = visuals["visuals"][0]
        self.assertTrue(page["observation_mode_selected"])
        self.assertTrue(page["observation_mode_overridden"])
        self.assertEqual(
            page["analysis_purpose"],
            "product_fact_candidates",
        )
        self.assertFalse(
            page["qwen_call_avoided_by_text_layer_context"]
        )
        self.assertEqual(summary["mixed_pdf"]["pages_selected_for_observations"], 1)
        self.assertEqual(summary["mixed_pdf"]["pages_processed_as_observations"], 0)


if __name__ == "__main__":
    unittest.main()
