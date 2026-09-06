from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

from product_evidence_guard.confirmation import apply_decision, ConfirmationRequestError
from product_evidence_guard.engine import analyze_directory
from product_evidence_guard.content_check import check_draft
from product_evidence_guard import workflow


class ContentWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "input"
        self.output = self.root / "output"
        self.source.mkdir()
        (self.source / "private-supplier.txt").write_text("净重：320g\n型号：PEG-01\n输出电压：5V\n", encoding="utf-8")
        self.summary = analyze_directory(self.source, self.output)
        self.product = json.loads((self.output / "product-facts.json").read_text(encoding="utf-8"))
        self.candidates = self.product["candidates"]

    def confirm(self, field="net_weight"):
        candidate = next(c for c in self.candidates if c["field"] == field)
        apply_decision(self.output, session_id=self.summary["session_id"], candidate_id=candidate["candidate_id"],
                       action="confirm", reason="私密供应商依据，不得发送到云端")
        return candidate

    def grant(self, fields=("net_weight",)):
        ids = [self.confirm(f)["candidate_id"] for f in fields]
        result = workflow.create_handoff(self.output, candidate_ids=ids, recipient="Qoder", purpose="生成商品介绍并回检")
        self.bundle = result["bundle_id"]
        return workflow.get_handoff(self.output, bundle_id=self.bundle, recipient="Qoder")

    def check(self, text):
        path = self.root / "draft.md"
        path.write_text(text, encoding="utf-8")
        return workflow.check_content(self.output, content_file=str(path), bundle_id=self.bundle, recipient="Qoder")

    def test_unconfirmed_and_empty_selection_cannot_be_authorized(self):
        for ids in ([], [self.candidates[0]["candidate_id"]]):
            with self.assertRaises(ConfirmationRequestError):
                workflow.create_handoff(self.output, candidate_ids=ids, recipient="Qoder", purpose="介绍")

    def test_packet_contains_only_selected_facts_and_no_evidence(self):
        packet = self.grant()
        serialized = json.dumps(packet, ensure_ascii=False)
        for forbidden in ("private-supplier", "私密", "confirmation_reason", "source_file", "locator", "raw_text", "PEG-01"):
            self.assertNotIn(forbidden, serialized)
        self.assertEqual(len(packet["facts"]), 1)
        self.assertEqual(packet["facts"][0]["value"], 320)
        with self.assertRaises(ConfirmationRequestError):
            workflow.get_handoff(self.output, bundle_id=self.bundle, recipient="Other")

    def test_conversion_and_conflict_then_revision(self):
        self.grant()
        first = self.check("净重：350g")
        self.assertEqual(first["status"], "blocked")
        self.assertEqual(first["findings"][0]["expected"][0]["value"], 320)
        second = self.check("净重：0.32kg")
        self.assertEqual(second["status"], "covered_fields_match")
        self.assertEqual(second["revision"], 2)
        self.assertEqual(first["artifact_id"], second["artifact_id"])
        self.assertEqual(len(workflow._manifest(self.output, self.summary["session_id"])["deliverables"]), 1)

    def test_unapproved_field_is_not_used_even_if_confirmed(self):
        self.confirm("model")
        self.grant()
        result = self.check("净重：320g。型号：PEG-01")
        self.assertEqual(result["status"], "blocked")
        self.assertTrue(any(f["kind"] == "unsupported" for f in result["findings"]))

    def test_negative_and_unlabeled_claims_are_not_silently_passed(self):
        self.grant()
        self.assertEqual(self.check("净重：-320g")["status"], "blocked")
        self.assertEqual(self.check("净重320g。轻至350g")["status"], "needs_review")
        self.assertEqual(self.check("外观时尚，美观大方。")["status"], "needs_review")

    def test_electrical_scopes_are_preserved(self):
        self.grant(("voltage",))
        self.assertEqual(self.check("输出电压：5V")["status"], "covered_fields_match")
        self.assertEqual(self.check("输入电压：5V")["status"], "needs_review")

    def test_source_change_invalidates_packet_and_deliverable_without_reanalysis(self):
        self.grant()
        self.check("产品介绍\n净重：320g")
        (self.source / "private-supplier.txt").write_text("净重：330g", encoding="utf-8")
        with self.assertRaises(ConfirmationRequestError):
            workflow.get_handoff(self.output, bundle_id=self.bundle, recipient="Qoder")
        snapshot = workflow.task_snapshot(self.output, detailed=True)
        self.assertEqual(snapshot["deliverables"][0]["status"], "source_stale")
        self.assertEqual(snapshot["deliverables"][0]["affected_references"][0]["lines"], [2])

    def test_changed_content_requires_recheck(self):
        self.grant()
        self.check("净重：320g")
        (self.root / "draft.md").write_text("净重：500g", encoding="utf-8")
        self.assertEqual(workflow.task_snapshot(self.output, detailed=True)["deliverables"][0]["status"], "content_changed")

    def test_revoked_grant_stops_future_reads(self):
        self.grant()
        workflow.revoke_handoff(self.output, bundle_id=self.bundle)
        with self.assertRaises(ConfirmationRequestError):
            workflow.get_handoff(self.output, bundle_id=self.bundle, recipient="Qoder")

    def test_modified_packet_is_rejected(self):
        self.grant()
        path = self.output / workflow.WORKFLOW_FILE
        data = json.loads(path.read_text(encoding="utf-8"))
        packet = data["bundles"][self.bundle]["packet"]
        packet["facts"][0]["value"] = 999
        packet["version"] = workflow.digest({k: v for k, v in packet.items() if k != "version"})
        path.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaises(ConfirmationRequestError):
            workflow.get_handoff(self.output, bundle_id=self.bundle, recipient="Qoder")

    def test_brief_task_does_not_include_private_source(self):
        self.grant()
        snapshot = workflow.task_snapshot(self.output)
        self.assertIsInstance(snapshot["performance"]["analysis_seconds"], (int, float))
        brief = json.dumps(snapshot, ensure_ascii=False)
        self.assertNotIn("private-supplier", brief)
        self.assertNotIn("私密", brief)
        self.assertNotIn('"value":', brief)
        self.assertNotIn('"normalized_value":', brief)
        metrics = workflow.performance_summary({"analysis_seconds": 1.25, "model_load_seconds": 0,
            "local_ai": {"model_reused": True, "model_path": "private-supplier/model"},
            "unchanged_files_reused": ["private-supplier.txt"]})
        self.assertEqual(metrics, {"analysis_seconds": 1.25, "model_load_seconds": 0,
                                  "model_reused": True, "unchanged_files_reused": 1})

    def test_offline_table_has_version_references(self):
        self.confirm()
        result = workflow.export_table(self.output)
        self.assertIn("320", Path(result["path"]).read_text(encoding="utf-8-sig"))
        detailed = workflow.task_snapshot(self.output, detailed=True)
        self.assertEqual(detailed["deliverables"][0]["status"], "covered_fields_match")
        self.assertEqual(len(detailed["bundles"]), 0)
        (self.source / "private-supplier.txt").write_text("净重：400g", encoding="utf-8")
        self.assertEqual(workflow.task_snapshot(self.output, detailed=True)["deliverables"][0]["status"], "source_stale")

    def test_excel_review_tracks_decisions_sources_and_keeps_strings_literal(self):
        from openpyxl import load_workbook
        from product_evidence_guard.review_workbook import export_review
        (self.source / "formula.txt").write_text("材质：=1+1\n", encoding="utf-8")
        self.summary = analyze_directory(self.source, self.output)
        self.candidates = json.loads((self.output / "product-facts.json").read_text("utf-8"))["candidates"]
        exported = export_review(self.output)
        self.assertFalse(exported["model_called"])
        book = load_workbook(exported["path"])
        try:
            self.assertEqual(book.sheetnames, ["核验总览", "候选与证据", "已确认参数"])
            self.assertEqual(book["候选与证据"].freeze_panes, "C5")
            self.assertEqual(book["已确认参数"].max_row, 4)
            self.assertIn("=1+1", [c.value for row in book["候选与证据"].iter_rows(min_row=5) for c in row])
            self.assertFalse(any(c.data_type == "f" for s in book for row in s for c in row))
        finally:
            book.close()
        self.confirm()
        report = workflow.task_snapshot(self.output)["deliverable_updates"][0]
        self.assertEqual((report["status"], report["next_action"]), ("needs_refresh", "export_review"))
        self.assertEqual(export_review(self.output)["confirmed_row_count"], 1)
        (self.source / "private-supplier.txt").write_text("净重：330g", encoding="utf-8")
        self.assertEqual(workflow.task_snapshot(self.output)["deliverable_updates"][0]["status"], "needs_refresh")
        stale = export_review(self.output)
        self.assertTrue(stale["needs_reanalysis"])
        self.assertEqual(stale["confirmed_row_count"], 0)

    def test_excel_does_not_present_conflicting_confirmed_values_as_usable(self):
        from product_evidence_guard.review_workbook import export_review
        self.confirm()
        (self.source / "other.txt").write_text("净重：300g", encoding="utf-8")
        summary = analyze_directory(self.source, self.output)
        candidates = json.loads((self.output / "product-facts.json").read_text("utf-8"))["candidates"]
        for candidate in candidates:
            if candidate["field"] == "net_weight":
                apply_decision(self.output, session_id=summary["session_id"], candidate_id=candidate["candidate_id"],
                               action="confirm", reason="测试冲突确认不能直接交付")
        self.assertEqual(export_review(self.output)["confirmed_row_count"], 0)
        with self.assertRaises(ConfirmationRequestError):
            workflow.export_table(self.output)

    def test_quantity_wrong_unit_is_not_accepted(self):
        facts = [{"fact_id": "quantity-fact", "field": "quantity", "value": 2, "unit": "count", "scope": None}]
        self.assertEqual(check_draft("数量：2件", facts)["status"], "covered_fields_match")
        self.assertEqual(check_draft("数量：2kg", facts)["status"], "needs_review")

    def test_dimensions_require_consistent_units(self):
        facts = [{"fact_id": "dimensions-fact", "field": "dimensions", "value": [100, 200, 300], "unit": "mm", "scope": None}]
        self.assertEqual(check_draft("尺寸：10×20×30cm", facts)["status"], "covered_fields_match")
        self.assertEqual(check_draft("尺寸：10×20×30mm", facts)["status"], "blocked")

    def test_empty_extraction_is_not_ready_for_export(self):
        (self.source / "private-supplier.txt").write_text("一段没有商品参数的介绍。", encoding="utf-8")
        analyze_directory(self.source, self.output)
        self.assertEqual(workflow.task_snapshot(self.output)["phase"], "no_candidates")
        with self.assertRaises(ConfirmationRequestError):
            workflow.export_table(self.output)

    def test_local_delivery_needs_no_cloud_grant_and_tracks_source_updates(self):
        self.confirm()
        result = workflow.export_local(self.output, reason="quota_exceeded")
        self.assertFalse(result["network_sent"])
        self.assertFalse(result["model_called"])
        self.assertEqual(result["status"], "local_delivery_ready")
        brief = Path(result["artifacts"]["brief"]).read_text(encoding="utf-8")
        self.assertIn("320", brief)
        self.assertNotIn("private-supplier", brief)
        self.assertNotIn("PEG-01", brief)
        snapshot = workflow.task_snapshot(self.output, detailed=True)
        self.assertEqual(snapshot["local_delivery"]["status"], "ready")
        self.assertEqual(snapshot["next_action"], "local_delivery_ready")
        self.assertEqual(len(snapshot["bundles"]), 0)
        self.assertEqual(len(snapshot["deliverables"]), 2)
        (self.source / "private-supplier.txt").write_text("净重：330g", encoding="utf-8")
        changed = workflow.task_snapshot(self.output, detailed=True)
        self.assertEqual(changed["local_delivery"]["status"], "needs_refresh")
        self.assertTrue(all(d["status"] == "source_stale" for d in changed["deliverables"]))

    def test_local_delivery_does_not_select_between_conflicting_confirmed_values(self):
        self.confirm()
        (self.source / "another.txt").write_text("净重：300g", encoding="utf-8")
        summary = analyze_directory(self.source, self.output)
        data = json.loads((self.output / "product-facts.json").read_text(encoding="utf-8"))
        other = next(c for c in data["candidates"] if c["normalized_value"] == 300)
        apply_decision(self.output, session_id=summary["session_id"], candidate_id=other["candidate_id"],
                       action="confirm", reason="测试中显式确认另一条矛盾值")
        with self.assertRaises(ConfirmationRequestError):
            workflow.export_local(self.output)
        self.assertFalse((self.output / "local-product-brief.md").exists())


if __name__ == "__main__":
    unittest.main()
