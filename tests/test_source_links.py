from __future__ import annotations

import hashlib
import io
from html import escape
import json
from pathlib import Path, PureWindowsPath
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import urlsplit
from urllib.request import url2pathname

from product_evidence_guard import conversation, review_assets, workflow
from product_evidence_guard.engine import analyze_directory
from product_evidence_guard.confirmation import ConfirmationRequestError
from product_evidence_guard.review_server import ReviewServer
from product_evidence_guard.review_workbook import export_review
from product_evidence_guard.source_links import write_source_links
from tests.test_fact_status import candidate


def local_path(url):
    return Path(url2pathname(urlsplit(url).path))


class SourceLinkTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source, self.output = self.root / "原始资料", self.root / "output"
        self.source.mkdir()
        (self.source / "主参数 (1).txt").write_text("净重：320g\n电压：12V\n型号：PRIVATE-42\n", encoding="utf-8")
        (self.source / "补充 #2.txt").write_text("净重：300g\n", encoding="utf-8")
        self.analysis = analyze_directory(self.source, self.output)

    def summary(self):
        permission = conversation.allow_review(self.output, recipient="WorkBuddy",
            fields=["net_weight", "voltage"], reason="展示冲突和待确认的原文链接")
        return conversation.review_summary(self.output, review_id=permission["review_id"], recipient="WorkBuddy")

    def test_each_conflict_and_pending_choice_has_matching_private_local_link(self):
        summary = self.summary()
        groups = summary["groups"]
        self.assertEqual({g["fact_status"] for g in groups}, {"conflict", "pending_confirmation"})
        for group in groups:
            for choice in group["choices"]:
                page = local_path(choice["source_url"])
                self.assertTrue(page.resolve().is_relative_to(self.output.resolve()))
                html = page.read_text("utf-8")
                self.assertIn(str(choice["value"]), html)
                self.assertIn("打开来源文件", html)
                self.assertIn("行：", html)
                self.assertNotIn("PRIVATE-42", html)
        serialized = json.dumps(summary, ensure_ascii=False)
        for secret in ("PRIVATE-42", "主参数 (1).txt", "补充 #2.txt", "raw_text", "candidate_id"):
            self.assertNotIn(secret, serialized)

    def test_static_reports_put_links_on_main_results_not_only_hidden_details(self):
        html = (self.output / "evidence-report.html").read_text("utf-8")
        markdown = (self.output / "conflicts.md").read_text("utf-8")
        self.assertIn("查看原文</a>", html.split("<details>")[0])
        self.assertIn("[查看原文](<file:", markdown.split("<details>")[0])
        self.assertNotIn("javascript:", html)

    def test_excel_keeps_one_fact_row_and_real_hyperlinks_for_every_option(self):
        from openpyxl import load_workbook
        result = export_review(self.output)
        book = load_workbook(result["path"])
        self.addCleanup(book.close)
        self.assertEqual(len(book.sheetnames), 5)
        self.assertEqual(book["核验总览"].max_row, result["fact_count"] + 4)
        counts = {}
        for name in ("冲突", "待确认事实"):
            sheet = book[name]
            self.assertEqual(sheet.max_row, 5)
            links = [c for row in sheet.iter_rows(min_row=5) for c in row if c.hyperlink]
            counts[name] = len(links)
            for cell in links:
                self.assertIn("查看原文", cell.value)
                self.assertTrue(local_path(cell.hyperlink.target).is_file())
                self.assertIn("行", cell.hyperlink.tooltip)
        self.assertEqual(counts, {"冲突": 2, "待确认事实": 1})
        self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))

    def test_changed_and_missing_sources_disable_original_file_link(self):
        before = self.summary()
        urls = [c["source_url"] for g in before["groups"] for c in g["choices"]]
        (self.source / "主参数 (1).txt").write_text("净重：999g\n", encoding="utf-8")
        (self.source / "补充 #2.txt").unlink()
        after = self.summary()
        self.assertTrue(after["needs_reanalysis"])
        for g in after["groups"]:
            for choice in g["choices"]:
                self.assertNotIn("value", choice)
                self.assertIn(choice["source_url"], urls)
                page = local_path(choice["source_url"]).read_text("utf-8")
                self.assertIn("来源已变化", page)
                self.assertNotIn("打开来源文件</a>", page)

    def test_pdf_page_link_and_escaped_snapshot_are_literal(self):
        source = self.source / '扫描 # (2).pdf'
        source.write_bytes(b"synthetic-file-for-link-test")
        c = candidate(source_file=source.name, locator={"page": 2},
                      file_hash=hashlib.sha256(source.read_bytes()).hexdigest(),
                      raw_text='<script>alert("x")</script> & 净重 320g')
        url = write_source_links(self.output, [c], input_root=self.source)[c.candidate_id]
        html = local_path(url).read_text("utf-8")
        self.assertIn(escape(source.as_uri() + "#page=2", quote=True), html)
        self.assertIn("页码：2", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)

    def test_cross_page_condition_link_uses_validated_same_file(self):
        source = self.source / "notes.pdf"
        source.write_bytes(b"synthetic-file-for-link-test")
        c = candidate(source_file=source.name, locator={"page": 2},
                      file_hash=hashlib.sha256(source.read_bytes()).hexdigest())
        c.provenance["source_conditions"] = [{"marker": "3<script>", "text": "Condition", "locator": {"page": 4}}]
        url = write_source_links(self.output, [c], input_root=self.source)[c.candidate_id]
        html = local_path(url).read_text("utf-8")
        self.assertIn(source.as_uri() + "#page=4", html)
        self.assertIn("查看附注 3&lt;script&gt;", html)
        self.assertNotIn("<script>", html)
        source.write_bytes(b"changed")
        url = write_source_links(self.output, [c], input_root=self.source)[c.candidate_id]
        self.assertNotIn("#page=4", local_path(url).read_text("utf-8"))

    def test_condition_preview_uses_stored_page_and_rechecks_source(self):
        from pypdf import PdfWriter
        from PIL import Image
        source = self.source / "notes.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=600)
        writer.add_blank_page(width=600, height=300)
        writer.write(source)
        c = {"candidate_id": "note-test", "source_file": source.name,
             "file_hash": hashlib.sha256(source.read_bytes()).hexdigest(), "locator": {"page": 1},
             "provenance": {"source_conditions": [{"marker": "1", "text": "Condition", "locator": {"page": 2}}]}}
        with patch("product_evidence_guard.review_server.workflow.context", return_value=(None, {"candidates": [c]}, None)):
            data, mime, _ = ReviewServer.preview(self.output, "note-test", condition="0")
            self.assertEqual(mime, "image/jpeg")
            with Image.open(io.BytesIO(data)) as image:
                self.assertGreater(image.width, image.height)
            for bad_index in ("-1", "1", "../2", "1e3"):
                with self.assertRaises(ConfirmationRequestError):
                    ReviewServer.preview(self.output, "note-test", condition=bad_index)
            source.write_bytes(b"changed")
            with self.assertRaises(ConfirmationRequestError):
                ReviewServer.preview(self.output, "note-test", condition="0")

    def test_unsafe_source_paths_and_executables_never_become_links(self):
        # Generate the synthetic absolute-path attack at runtime; the release
        # source itself must not carry machine-looking absolute paths.
        absolute_fixture = str(PureWindowsPath("C:", "/", "outside.txt"))
        self.assertTrue(PureWindowsPath(absolute_fixture).is_absolute())
        for name in ("../outside.txt", absolute_fixture, "\\\\server\\share\\x.txt", "bad.exe"):
            with self.subTest(name=name):
                c = candidate(source_file=name)
                html = local_path(write_source_links(self.output, [c], input_root=self.source)[c.candidate_id]).read_text("utf-8")
                self.assertNotIn("打开来源文件</a>", html)

    def test_link_output_refuses_symlink_redirection(self):
        folder = self.output / "source-links"
        outside = self.root / "outside"
        outside.mkdir()
        try:
            (folder / "redirect").symlink_to(outside, target_is_directory=True)
        except OSError:
            self.skipTest("Symlinks unavailable")
        with self.assertRaises(ValueError):
            write_source_links(folder / "redirect", [candidate()], input_root=self.source)

    def test_live_review_exposes_source_links_for_both_states_and_all_sources(self):
        self.assertIn("if(g.fact_status!=='verified'){const actions", review_assets.JS)
        self.assertIn("sources.forEach((s,i)=>", review_assets.JS)
        self.assertIn("function sourceLink(c,text='查看原文',condition=null)", review_assets.JS)
        self.assertIn("preview(c,condition).catch", review_assets.JS)

    @unittest.skipUnless(shutil.which("node"), "Node is unavailable for the renderer test")
    def test_live_option_renderer(self):
        subprocess.run([shutil.which("node"), str(Path(__file__).with_name("check_review_source_links.cjs"))],
                       check=True, capture_output=True, text=True, timeout=20)


if __name__ == "__main__":
    unittest.main()
