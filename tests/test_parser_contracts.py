"""No evidence is lost before ownership, value or status decisions are made."""
import csv
import json
from pathlib import Path
import tempfile
import unittest

from product_evidence_guard.content_check import check_draft
from product_evidence_guard.engine import analyze_directory
from product_evidence_guard.extractor import extract_rule_candidates
from product_evidence_guard.graph import build_graph
from product_evidence_guard.models import SourceBlock
from product_evidence_guard.normalization import normalize_value
from product_evidence_guard.structured_rows import bind_table
from product_evidence_guard.confirmation import apply_review_action


def extract(text, index=0, **provenance):
    return extract_rule_candidates(SourceBlock(str(index), "a.txt", "text", "hash",
        {"line": index + 1}, text, provenance=provenance))


class WholeValueTests(unittest.TestCase):
    def test_supported_qualifiers_are_not_scalar_values(self):
        for amount in (80, 275, 1200):
            for phrase, operator in (("不能超过", "≤"), ("不得超过", "≤"), ("不超过", "≤"),
                                     ("约为", "≈"), ("大约为", "≈"), ("approximately ", "≈")):
                with self.subTest(amount=amount, phrase=phrase):
                    raw = f"{phrase}{amount}g"
                    self.assertEqual(normalize_value("net_weight", raw).value, f"{operator}{amount}")
                    rows = extract("净重：" + raw)
                    self.assertEqual(rows[0].normalized_value, f"{operator}{amount}")
                    self.assertEqual(build_graph(rows)[1][0].fact_status, "verified")

    def test_unknown_semantics_never_disappear(self):
        for field, value in (("net_weight", "80g"), ("voltage", "5V"),
                             ("temperature", "0~40°C"), ("dimensions", "10x20mm")):
            for raw in ("非" + value, "not " + value, "除外 " + value, "未达到" + value,
                        "zqx " + value, value + " zqx", "大概可能" + value):
                with self.subTest(raw=raw):
                    result = normalize_value(field, raw)
                    self.assertIsNone(result.unit)
                    self.assertTrue(any(n.startswith("unparsed_") for n in result.notes))

    def test_draft_checks_the_full_qualified_expression(self):
        for raw in ("非80g", "not 80 grams", "不能超过80g", "约为80g", "80g excluded"):
            facts = [{"field": "net_weight", "value": 80, "unit": "g", "scope": "net"}]
            self.assertNotEqual(check_draft("净重：" + raw, facts)["status"], "covered_fields_match", raw)
        facts = [{"field": "net_weight", "value": "≤80", "unit": "g", "scope": "net"}]
        self.assertEqual(check_draft("净重：不能超过80g", facts)["status"], "covered_fields_match")

    def test_multi_quantity_expression_cannot_discard_negation(self):
        for raw in ("Input: not 5V / 2A", "输入：非5V / 2A", "Output: 5V / 2A excluded"):
            with self.subTest(raw=raw):
                rows = extract(raw)
                self.assertTrue(rows)
                self.assertNotIn("verified", {g.fact_status for g in build_graph(rows)[1]})

    def test_words_before_the_claim_label_are_not_discarded_either(self):
        facts = [{"field": "net_weight", "value": 80, "unit": "g", "scope": "net"}]
        for text in ("并非净重80g", "不是净重80g", "含附件时净重80g", "zqx 净重80g"):
            self.assertNotEqual(check_draft(text, facts)["status"], "covered_fields_match", text)
        facts = [{"field": "ip_rating", "value": "ip65", "unit": None, "scope": None}]
        for text in ("not IP65", "非IP65", "IP65 excluded"):
            self.assertNotEqual(check_draft(text, facts)["status"], "covered_fields_match", text)


class BindingTests(unittest.TestCase):
    def test_value_first_labels_are_case_invariant(self):
        for label, raw in (("ROM", "384 kB"), ("GPIO", "21"), ("ERRORS", "21"), ("NOVEL_COUNTER", "14")):
            signatures = []
            for spelling in (label, label.title(), label.lower()):
                rows = extract(raw + " " + spelling)
                self.assertEqual(len(rows), 1)
                signatures.append((rows[0].field, rows[0].normalized_value, build_graph(rows)[1][0].fact_status))
            self.assertEqual(len(set(signatures)), 1)
            self.assertEqual(signatures[0][2], "pending_confirmation")

    def test_explicit_binding_establishes_meaning_not_a_vote_for_the_value(self):
        for other, status in (("21", "verified"), ("22", "conflict")):
            for label in ("GPIO", "Gpio", "gpio"):
                rows = extract("21 " + label) + extract("gpio: " + other, 1)
                self.assertEqual(build_graph(rows)[1][0].fact_status, status)

    def test_matrix_keeps_every_nonstructural_column(self):
        rows = [["SKU", "Parameter", "Value", "Unit", "Unfamiliar operating context"],
                ["A-1", "净重", "80", "g", "Eco"]]
        bound = bind_table([(i, list(enumerate(row))) for i, row in enumerate(rows)], table_id="t")
        self.assertEqual(bound[0].identity.get("sku"), "A-1")
        self.assertIn("Eco", " ".join(bound[0].parent_labels))

    def test_repetition_of_ambiguous_labels_does_not_establish_meaning(self):
        rows = extract("21 errors", 0) + extract("21 ERRORS", 1)
        self.assertEqual(build_graph(rows)[1][0].fact_status, "pending_confirmation")
        rows[0].status = "confirmed"
        # A human decision is also a field binding for equal copies.
        self.assertEqual(build_graph(rows)[1][0].fact_status, "verified")


class EndToEndBindingTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root, self.output = Path(temp.name) / "input", Path(temp.name) / "output"
        self.root.mkdir()

    def run_rows(self, rows, extension="csv"):
        path = self.root / ("table." + extension)
        if extension == "csv":
            with path.open("w", encoding="utf8", newline="") as stream:
                csv.writer(stream).writerows(rows)
        elif extension == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            for row in rows:
                workbook.active.append(row)
            workbook.save(path)
            workbook.close()
        else:
            from docx import Document
            document = Document()
            table = document.add_table(rows=0, cols=len(rows[0]))
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
            document.save(path)
        summary = analyze_directory(self.root, self.output, preprocessing_workers=1)
        self.assertFalse(summary["errors"])
        data = json.loads((self.output / "product-facts.json").read_text("utf8"))
        return data, summary

    def test_matrix_ownership_in_all_office_readers(self):
        for extension in ("csv", "xlsx", "docx"):
            with self.subTest(extension=extension):
                data, _ = self.run_rows([["SKU", "Parameter", "Value", "Unit"],
                    ["PROD-A", "净重", "80", "g"], ["PROD-B", "净重", "90", "g"]], extension)
                groups = [g for g in data["facts"] if g["field"] == "net_weight"]
                self.assertEqual(len(groups), 2)
                self.assertEqual({g["fact_status"] for g in groups}, {"verified"})
                self.assertEqual(len({g["product_id"] for g in groups}), 2)
                current_format = [c for c in data["candidates"] if c["source_file"].endswith("." + extension)]
                self.assertEqual(len(current_format), 2)
                self.assertEqual({c["product_sku"] for c in current_format}, {"PROD-A", "PROD-B"})

    def test_context_column_renaming_never_creates_a_conflict(self):
        for header in ("Conditions", "Operating conditions", "适用条件", "实验场景甲"):
            with self.subTest(header=header):
                data, _ = self.run_rows([["Parameter", "Value", "Unit", header],
                    ["输入电压", "5", "V", "Eco"], ["输入电压", "12", "V", "Turbo"]])
                groups = [g for g in data["facts"] if g["field"] == "voltage"]
                self.assertEqual(len(groups), 2)
                self.assertEqual({g["fact_status"] for g in groups}, {"verified"})
                self.assertEqual({g["selected_value"] for g in groups}, {5, 12})
                self.assertTrue(all("input" in g["scope"].split("|") for g in groups))
                self.assertTrue(all("Eco" in c["raw_text"] or "Turbo" in c["raw_text"] for c in data["candidates"]))

    def test_missing_row_owner_is_not_assigned_to_another_product(self):
        data, _ = self.run_rows([["SKU", "Parameter", "Value", "Unit"],
            ["A-1", "净重", "80", "g"], ["", "净重", "90", "g"]])
        self.assertNotIn("conflict", {g["fact_status"] for g in data["facts"]})
        unknown = next(c for c in data["candidates"] if c["normalized_value"] == 90)
        self.assertIn(unknown["product_identity_status"], {"ambiguous", "unresolved"})

    def test_zero_value_and_zero_condition_are_not_blank_cells(self):
        data, _ = self.run_rows([["Parameter", "Value", "Unit", "Conditions"],
            ["工作温度", 0, "°C", 0]], "xlsx")
        self.assertEqual(len(data["facts"]), 1)
        self.assertEqual(data["facts"][0]["selected_value"], 0)
        self.assertIn("context:0", data["facts"][0]["scope"])

    def test_conflicting_identity_columns_are_not_guessed(self):
        data, _ = self.run_rows([["SKU", "sku", "Parameter", "Value", "Unit"],
            ["A-1", "B-1", "净重", "80", "g"]])
        self.assertEqual(data["facts"][0]["review_reason_code"], "identity_ambiguous")

    def test_human_resolution_updates_both_summary_copies(self):
        data, _ = self.run_rows([["Parameter", "Value"], ["净重", "80g"], ["净重", "90g"]])
        group = data["facts"][0]
        apply_review_action(self.output, session_id=data["run_summary"]["session_id"], action="adopt",
            group_id=group["group_id"], expected_candidate_ids=group["candidate_ids"],
            candidate_id=group["candidate_ids"][0])
        data = json.loads((self.output / "product-facts.json").read_text("utf8"))
        summary = json.loads((self.output / "run-summary.json").read_text("utf8"))
        for counts in (summary, data["run_summary"]):
            self.assertEqual((counts["conflict_count"], counts["review_count"], counts["pass_count"]), (0, 0, 1))

    def test_summary_and_report_use_fact_status(self):
        data, summary = self.run_rows([["Parameter", "Value"], ["净重", "80g"]])
        self.assertEqual(summary["review_count"], 0)
        self.assertEqual(summary["pass_count"], 1)
        self.assertEqual(data["run_summary"]["review_count"], 0)
        report = (self.output / "conflicts.md").read_text("utf8")
        self.assertIn("## 净重 — 已确认", report)
        self.assertNotIn("严重程度：**review**", report)
        self.assertNotIn("待人工确认：1", report)


if __name__ == "__main__":
    unittest.main()
