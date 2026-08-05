from __future__ import annotations

import argparse
from collections.abc import Callable
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import zipfile


GENERATION_REVISION = "document-visual-samples-v1"
FIXED_TIME = datetime(2024, 1, 1, tzinfo=timezone.utc)
ZIP_TIME = (2024, 1, 1, 0, 0, 0)
FIXED_W3CDTF = b"2024-01-01T00:00:00Z"
_MODIFIED_PROPERTY_PATTERN = re.compile(
    rb"(<dcterms:modified\b[^>]*>).*?(</dcterms:modified>)",
    flags=re.DOTALL,
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _font(size: int, *, bold: bool = False):
    from PIL import ImageFont

    windows_root = Path(os.environ.get("WINDIR", str(Path(os.sep) / "Windows")))
    names = (
        ("arialbd.ttf", "DejaVuSans-Bold.ttf")
        if bold
        else ("arial.ttf", "DejaVuSans.ttf")
    )
    candidates = [
        windows_root / "Fonts" / names[0],
        Path("/usr/share/fonts/truetype/dejavu") / names[1],
    ]
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _draw_spec_card(path: Path) -> None:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (1600, 1000), "#f7f9fc")
    draw = ImageDraw.Draw(image)
    title_font = _font(58, bold=True)
    heading_font = _font(36, bold=True)
    body_font = _font(34)
    small_font = _font(26)

    draw.rounded_rectangle(
        (45, 40, 1555, 960),
        radius=30,
        fill="white",
        outline="#26374a",
        width=5,
    )
    draw.rectangle((45, 40, 1555, 155), fill="#173f73")
    draw.text(
        (90, 68),
        "VECTORDOCK X1  |  PRODUCT SPECIFICATION",
        fill="white",
        font=title_font,
    )

    rows = (
        ("NET WEIGHT", "320 g"),
        ("INPUT VOLTAGE", "12 V"),
        ("RATED POWER", "65 W"),
        ("DIMENSIONS", "120 x 80 x 28 mm"),
    )
    y = 205
    for label, value in rows:
        draw.rounded_rectangle(
            (90, y, 790, y + 105),
            radius=14,
            fill="#eef4fb",
            outline="#8ca7c4",
            width=3,
        )
        draw.text((120, y + 18), label, fill="#173f73", font=small_font)
        draw.text((470, y + 13), value, fill="#111827", font=heading_font)
        y += 125

    draw.text((900, 195), "MECHANICAL DRAWING", fill="#173f73", font=heading_font)
    draw.rounded_rectangle(
        (930, 290, 1450, 650),
        radius=20,
        fill="#f4f7fa",
        outline="#26374a",
        width=5,
    )
    draw.line((930, 700, 1450, 700), fill="#c43d3d", width=4)
    draw.line((930, 680, 930, 720), fill="#c43d3d", width=4)
    draw.line((1450, 680, 1450, 720), fill="#c43d3d", width=4)
    draw.text((1120, 715), "120 mm", fill="#c43d3d", font=body_font)
    draw.line((1510, 290, 1510, 650), fill="#c43d3d", width=4)
    draw.line((1490, 290, 1530, 290), fill="#c43d3d", width=4)
    draw.line((1490, 650, 1530, 650), fill="#c43d3d", width=4)
    draw.text((1360, 445), "80 mm", fill="#c43d3d", font=small_font)

    draw.text((930, 775), "POWER BY MODE (W)", fill="#173f73", font=small_font)
    bars = (("ECO", 30, "#74a9cf"), ("BALANCED", 45, "#2b8cbe"), ("TURBO", 65, "#045a8d"))
    x = 930
    for label, value, color in bars:
        height = value
        draw.rectangle((x, 920 - height, x + 115, 920), fill=color)
        draw.text((x + 25, 925), label, fill="#111827", font=small_font)
        draw.text((x + 35, 880 - height), f"{value} W", fill="#111827", font=small_font)
        x += 180
    image.save(path, format="PNG", optimize=False)
    image.close()


def _normalize_xlsx_core_properties(data: bytes) -> bytes:
    normalized, replacement_count = _MODIFIED_PROPERTY_PATTERN.subn(
        rb"\g<1>" + FIXED_W3CDTF + rb"\g<2>",
        data,
        count=1,
    )
    if replacement_count != 1:
        raise ValueError(
            "XLSX core properties did not contain one modified timestamp"
        )
    return normalized


def _normalize_zip(
    path: Path,
    *,
    member_transforms: dict[str, Callable[[bytes], bytes]] | None = None,
) -> None:
    transforms = member_transforms or {}
    with tempfile.NamedTemporaryFile(
        prefix=f"{path.stem}-",
        suffix=path.suffix,
        dir=path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temporary_path,
            "w",
        ) as target:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                normalized = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
                normalized.compress_type = info.compress_type
                normalized.comment = info.comment
                normalized.extra = b""
                normalized.create_system = info.create_system
                normalized.external_attr = info.external_attr
                normalized.internal_attr = info.internal_attr
                normalized.flag_bits = info.flag_bits
                data = source.read(info.filename)
                transform = transforms.get(info.filename)
                if transform is not None:
                    data = transform(data)
                target.writestr(normalized, data)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _write_docx(path: Path, image_path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.core_properties.title = "Controlled embedded image sample"
    document.core_properties.author = "Product Evidence Guard"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.add_heading("Embedded product drawing", level=1)
    document.add_paragraph(
        "The visual below is intentionally embedded inside this DOCX file."
    )
    shape = document.add_picture(str(image_path), width=Inches(6.4))
    shape._inline.docPr.set("descr", "VectorDock X1 specification and drawing")
    document.add_paragraph(
        "Expected behavior: preserve the document and paragraph location for every "
        "visual transcription."
    )
    document.save(path)
    _normalize_zip(path)


def _write_xlsx(path: Path, image_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as SpreadsheetImage

    workbook = Workbook()
    workbook.properties.creator = "Product Evidence Guard"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    worksheet = workbook.active
    worksheet.title = "Performance"
    worksheet.append(["Mode", "Power (W)"])
    worksheet.append(["Eco", 30])
    worksheet.append(["Balanced", 45])
    worksheet.append(["Turbo", 65])

    chart = BarChart()
    chart.title = "Power by mode"
    chart.y_axis.title = "W"
    chart.x_axis.title = "Mode"
    chart.add_data(
        Reference(worksheet, min_col=2, min_row=1, max_row=4),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(worksheet, min_col=1, min_row=2, max_row=4)
    )
    chart.height = 8
    chart.width = 13
    worksheet.add_chart(chart, "A7")

    embedded = SpreadsheetImage(str(image_path))
    embedded.width = 800
    embedded.height = 500
    worksheet.add_image(embedded, "D2")
    workbook.save(path)
    workbook.close()
    _normalize_zip(
        path,
        member_transforms={
            "docProps/core.xml": _normalize_xlsx_core_properties,
        },
    )


def _pdf_string(value: str) -> bytes:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    return escaped.encode("ascii")


def _write_mixed_pdf(path: Path, image_path: Path) -> None:
    from PIL import Image
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    intermediate = path.with_name(f".{path.name}.image-only.tmp.pdf")
    image = Image.open(image_path).convert("RGB")
    try:
        image.save(intermediate, format="PDF", resolution=150)
    finally:
        image.close()

    try:
        reader = PdfReader(intermediate)
        page = reader.pages[0]
        existing = page.get_contents()
        existing_data = existing.get_data() if existing is not None else b""

        resources = page["/Resources"].get_object()
        fonts = resources.get("/Font")
        if fonts is None:
            fonts = DictionaryObject()
            resources[NameObject("/Font")] = fonts
        fonts[NameObject("/FPEG")] = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        layer_text = (
            b"\nBT /FPEG 11 Tf 24 20 Td ("
            + _pdf_string(
                "Controlled mixed visual page - product specification diagram"
            )
            + b") Tj ET\n"
        )
        combined = DecodedStreamObject()
        combined.set_data(existing_data + layer_text)
        page[NameObject("/Contents")] = combined

        writer = PdfWriter()
        writer.add_page(page)
        writer.add_metadata(
            {
                "/Title": "Controlled mixed visual PDF sample",
                "/Author": "Product Evidence Guard",
                "/Producer": "Product Evidence Guard sample generator",
                "/CreationDate": "D:20240101000000Z",
                "/ModDate": "D:20240101000000Z",
            }
        )
        with path.open("wb") as handle:
            writer.write(handle)
    finally:
        intermediate.unlink(missing_ok=True)


def generate(output_root: Path) -> dict[str, object]:
    output_root = output_root.expanduser().resolve()
    input_dir = output_root / "input"
    assets_dir = output_root / "assets"
    input_dir.mkdir(parents=True, exist_ok=True)
    assets_dir.mkdir(parents=True, exist_ok=True)

    image_path = assets_dir / "vectordock-x1-spec-card.png"
    docx_path = input_dir / "controlled-embedded-image.docx"
    xlsx_path = input_dir / "controlled-chart-and-image.xlsx"
    pdf_path = input_dir / "controlled-mixed-visual.pdf"

    _draw_spec_card(image_path)
    _write_docx(docx_path, image_path)
    _write_xlsx(xlsx_path, image_path)
    _write_mixed_pdf(pdf_path, image_path)

    manifest: dict[str, object] = {
        "schema_version": 1,
        "generation_revision": GENERATION_REVISION,
        "input_directory": str(input_dir),
        "files": [
            {
                "path": path.relative_to(output_root).as_posix(),
                "sha256": _sha256(path),
                "size_bytes": path.stat().st_size,
            }
            for path in (docx_path, xlsx_path, pdf_path)
        ],
        "asset": {
            "path": image_path.relative_to(output_root).as_posix(),
            "sha256": _sha256(image_path),
            "size_bytes": image_path.stat().st_size,
        },
        "expectations": {
            "docx_embedded_images": 1,
            "xlsx_embedded_images": 1,
            "xlsx_native_charts": 1,
            "pdf_mixed_visual_pages": [1],
            "visual_facts": {
                "net_weight": "320 g",
                "input_voltage": "12 V",
                "rated_power": "65 W",
                "dimensions": "120 x 80 x 28 mm",
            },
            "native_chart_points": [
                {"category": "Eco", "value": 30},
                {"category": "Balanced", "value": 45},
                {"category": "Turbo", "value": 65},
            ],
        },
    }
    manifest_path = output_root / "expected-manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return {**manifest, "manifest": str(manifest_path)}


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Generate controlled DOCX, XLSX and mixed-PDF samples for the "
            "document visual pipeline."
        )
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(".runtime/document-visual-samples"),
        help="Output root (default: .runtime/document-visual-samples).",
    )
    args = parser.parse_args()
    result = generate(args.output)
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
