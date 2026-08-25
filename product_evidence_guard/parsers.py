from __future__ import annotations

import csv
from dataclasses import dataclass
import hashlib
from itertools import islice
import json
from math import ceil, isfinite
from pathlib import Path
import re
from typing import Any, Iterable
import zipfile

from .document_visuals import (
    EmbeddedVisualAsset,
    StructuredDocumentVisual,
    extract_docx_visuals,
    extract_xlsx_visuals,
)
from .models import SourceBlock


TEXT_EXTENSIONS = {".txt", ".md"}
DOCUMENT_EXTENSIONS = {".docx", ".xlsx", ".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".csv", ".json"} | DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
MAX_PDF_PAGES = 100
PDF_RENDER_SCALE = 2.0
MAX_PDF_RENDER_PIXELS = 25_000_000
MAX_PDF_MIXED_VISUAL_PAGES = 12
PDF_LARGE_IMAGE_MIN_PIXELS = 160_000
PDF_VECTOR_VISUAL_OPERATOR_THRESHOLD = 24
PDF_VECTOR_VISUAL_HIGH_THRESHOLD = 120
PDF_XOBJECT_MAX_DEPTH = 4
PDF_XOBJECT_MAX_OBJECTS = 128
PDF_VISUAL_ROI_MARGIN_POINTS = 36.0
PDF_VISUAL_ROI_TEXT_SNAP_POINTS = 48.0
PDF_VISUAL_ROI_MIN_SAVINGS_RATIO = 0.20
PDF_VISUAL_ROI_MAX_AREA_RATIO = 0.80
PDF_VISUAL_ROI_MIN_PLACED_AREA_RATIO = 0.005
PDF_VISUAL_ROI_MIN_PLACED_DIMENSION_RATIO = 0.03
PDF_VISUAL_ROI_HEADER_FOOTER_RATIO = 0.07

_PDF_VISUAL_KEYWORDS = (
    "chart",
    "curve",
    "diagram",
    "drawing",
    "dimension",
    "figure",
    "graph",
    "mechanical",
    "physical specification",
    "pinout",
    "plot",
    "schematic",
    "typical performance",
    "typical characteristics",
    "尺寸",
    "工程图",
    "机械图",
    "曲线",
    "示意图",
    "图纸",
)

_PDF_VISUAL_UNIT_PATTERN = re.compile(
    r"\d[\d.,]*\s*(?:mm|cm|inch|in|v(?:ac|dc)?|a|w|mah|wh|kg|g|oz|lb|°c)\b",
    flags=re.IGNORECASE,
)

_PDF_VECTOR_OPERATORS = {
    b"B",
    b"B*",
    b"F",
    b"S",
    b"b",
    b"b*",
    b"c",
    b"f",
    b"f*",
    b"l",
    b"m",
    b"re",
    b"s",
    b"v",
    b"y",
}


@dataclass(frozen=True, slots=True)
class DocumentParseResult:
    blocks: tuple[SourceBlock, ...]
    visual_assets: tuple[EmbeddedVisualAsset, ...] = ()
    structured_visuals: tuple[StructuredDocumentVisual, ...] = ()
    visual_issues: tuple[dict[str, Any], ...] = ()


@dataclass(frozen=True, slots=True)
class PdfParseResult:
    blocks: tuple[SourceBlock, ...]
    scanned_page_numbers: tuple[int, ...]
    page_count: int
    mixed_visual_page_numbers: tuple[int, ...] = ()
    visual_page_reasons: tuple[dict[str, Any], ...] = ()


@dataclass(frozen=True, slots=True)
class RenderedPdfPage:
    page_number: int
    image_path: Path
    pixel_width: int
    pixel_height: int


@dataclass(frozen=True, slots=True)
class RenderedPdfVisualRegion(RenderedPdfPage):
    """A rendered PDF page or conservative raster region from that page.

    ``crop_box_pdf`` uses PDF canvas coordinates in the order
    ``(left, bottom, right, top)``. A full-page fallback intentionally remains
    compatible with :class:`RenderedPdfPage` while making the crop decision
    auditable by callers.
    """

    crop_box_pdf: tuple[float, float, float, float]
    crop_ratio: float
    strategy: str
    fallback_reason: str | None
    page_width_pdf: float
    page_height_pdf: float


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _block_id(relative_path: str, file_hash: str, locator: dict[str, Any], text: str) -> str:
    payload = json.dumps(
        {"path": relative_path, "hash": file_hash, "locator": locator, "text": text},
        sort_keys=True,
        ensure_ascii=False,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()[:20]


def _make_block(
    *,
    relative_path: str,
    source_kind: str,
    file_hash: str,
    locator: dict[str, Any],
    text: str,
    confidence: float = 1.0,
    method: str = "deterministic_parser",
    confidence_source: str = "deterministic",
) -> SourceBlock:
    return SourceBlock(
        block_id=_block_id(relative_path, file_hash, locator, text),
        source_file=relative_path,
        source_kind=source_kind,
        file_hash=file_hash,
        locator=locator,
        text=text,
        recognition_confidence=max(0.0, min(1.0, confidence)),
        extraction_method=method,
        recognition_confidence_source=confidence_source,
        provenance={"recognition_confidence_source": confidence_source},
    )


def parse_text(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    text = path.read_text(encoding="utf-8-sig")
    blocks: list[SourceBlock] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        value = line.strip()
        if value:
            blocks.append(
                _make_block(
                    relative_path=relative_path,
                    source_kind="text",
                    file_hash=file_hash,
                    locator={"line": line_number},
                    text=value,
                )
            )
    return blocks


def parse_csv(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    blocks: list[SourceBlock] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for row_number, row in enumerate(reader, start=1):
            nonempty = [(index + 1, cell.strip()) for index, cell in enumerate(row) if cell.strip()]
            if not nonempty:
                continue
            if len(nonempty) >= 2:
                combined = " | ".join(cell for _, cell in nonempty)
                blocks.append(
                    _make_block(
                        relative_path=relative_path,
                        source_kind="csv_row",
                        file_hash=file_hash,
                        locator={"row": row_number, "columns": [index for index, _ in nonempty]},
                        text=combined,
                    )
                )
            else:
                column, value = nonempty[0]
                blocks.append(
                    _make_block(
                        relative_path=relative_path,
                        source_kind="csv_cell",
                        file_hash=file_hash,
                        locator={"row": row_number, "column": column},
                        text=value,
                    )
                )
    return blocks


def _flatten_json(value: Any, path: str = "$") -> Iterable[tuple[str, Any]]:
    if isinstance(value, dict):
        for key, item in value.items():
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            yield from _flatten_json(item, f"{path}/{escaped}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from _flatten_json(item, f"{path}/{index}")
    elif value is not None:
        yield path, value


def parse_json(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    blocks: list[SourceBlock] = []
    for json_path, value in _flatten_json(data):
        label = json_path.rsplit("/", 1)[-1].replace("~1", "/").replace("~0", "~")
        text = f"{label}: {value}"
        blocks.append(
            _make_block(
                relative_path=relative_path,
                source_kind="json_value",
                file_hash=file_hash,
                locator={"json_path": json_path},
                text=text,
            )
        )
    return blocks


def parse_ocr_sidecar(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    blocks_data = data.get("blocks")
    if not isinstance(blocks_data, list):
        raise ValueError(f"OCR sidecar blocks must be an array: {relative_path}")
    image_name = str(data.get("image") or relative_path.removesuffix(".ocr.json"))
    blocks: list[SourceBlock] = []
    for index, item in enumerate(blocks_data):
        if not isinstance(item, dict) or not str(item.get("text", "")).strip():
            continue
        locator = {"image": image_name, "block": index}
        if isinstance(item.get("bbox"), list):
            locator["bbox"] = item["bbox"]
        blocks.append(
            _make_block(
                relative_path=relative_path,
                source_kind="image_ocr",
                file_hash=file_hash,
                locator=locator,
                text=str(item["text"]).strip(),
                confidence=float(item.get("confidence", 0.75)),
                method=str(item.get("method", "ocr_sidecar")),
                confidence_source=str(
                    item.get("confidence_source", "ocr_sidecar")
                ),
            )
        )
    return blocks


def parse_docx_document(
    path: Path,
    relative_path: str,
    file_hash: str,
) -> DocumentParseResult:
    try:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("读取 DOCX 需要安装 python-docx：pip install -e '.[documents]'") from exc

    document = Document(path)
    blocks: list[SourceBlock] = []
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            text = Paragraph(child, document).text.strip()
            if text:
                blocks.append(
                    _make_block(
                        relative_path=relative_path,
                        source_kind="docx_paragraph",
                        file_hash=file_hash,
                        locator={"paragraph": paragraph_index},
                        text=text,
                    )
                )
            paragraph_index += 1
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            for row_index, row in enumerate(table.rows):
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    blocks.append(
                        _make_block(
                            relative_path=relative_path,
                            source_kind="docx_table_row",
                            file_hash=file_hash,
                            locator={"table": table_index, "row": row_index},
                            text=" | ".join(values),
                        )
                    )
            table_index += 1
    visual_assets, structured_visuals, visual_issues = extract_docx_visuals(
        document,
        relative_path=relative_path,
        file_hash=file_hash,
    )
    return DocumentParseResult(
        blocks=tuple(blocks),
        visual_assets=tuple(visual_assets),
        structured_visuals=tuple(structured_visuals),
        visual_issues=tuple(visual_issues),
    )


def parse_docx(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    return list(parse_docx_document(path, relative_path, file_hash).blocks)


def _xlsx_has_visual_parts(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            return any(
                name.startswith(("xl/charts/", "xl/drawings/", "xl/media/"))
                for name in archive.namelist()
            )
    except (OSError, zipfile.BadZipFile):
        return False


def parse_xlsx_document(
    path: Path,
    relative_path: str,
    file_hash: str,
) -> DocumentParseResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("读取 XLSX 需要安装 openpyxl：pip install -e '.[documents]'") from exc

    workbook = load_workbook(path, read_only=True, data_only=False)
    blocks: list[SourceBlock] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                cells = [(cell.coordinate, cell.value) for cell in row if cell.value not in (None, "")]
                if not cells:
                    continue
                text = " | ".join(str(value) for _, value in cells)
                blocks.append(
                    _make_block(
                        relative_path=relative_path,
                        source_kind="xlsx_row",
                        file_hash=file_hash,
                        locator={"sheet": sheet.title, "cells": [coordinate for coordinate, _ in cells]},
                        text=text,
                    )
                )
    finally:
        workbook.close()

    visual_assets: list[EmbeddedVisualAsset] = []
    structured_visuals: list[StructuredDocumentVisual] = []
    visual_issues: list[dict[str, Any]] = []
    if _xlsx_has_visual_parts(path):
        visual_workbook = load_workbook(
            path,
            read_only=False,
            data_only=False,
        )
        try:
            (
                visual_assets,
                structured_visuals,
                visual_issues,
            ) = extract_xlsx_visuals(
                visual_workbook,
                relative_path=relative_path,
                file_hash=file_hash,
            )
        finally:
            visual_workbook.close()
    return DocumentParseResult(
        blocks=tuple(blocks),
        visual_assets=tuple(visual_assets),
        structured_visuals=tuple(structured_visuals),
        visual_issues=tuple(visual_issues),
    )


def parse_xlsx(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    return list(parse_xlsx_document(path, relative_path, file_hash).blocks)


def _resolve_pdf_object(value: Any) -> Any:
    try:
        return value.get_object()
    except (AttributeError, KeyError, TypeError, ValueError):
        return value


def _pdf_large_image_count(page: Any) -> int:
    try:
        resources = _resolve_pdf_object(page.get("/Resources"))
        xobjects = _resolve_pdf_object(resources.get("/XObject"))
    except (AttributeError, KeyError, TypeError, ValueError):
        return 0

    seen: set[tuple[Any, ...]] = set()
    large_images = 0

    def visit(raw_xobjects: Any, depth: int) -> None:
        nonlocal large_images
        if (
            depth > PDF_XOBJECT_MAX_DEPTH
            or len(seen) >= PDF_XOBJECT_MAX_OBJECTS
        ):
            return
        try:
            resolved_xobjects = _resolve_pdf_object(raw_xobjects)
            values = list(resolved_xobjects.values())
        except (AttributeError, KeyError, TypeError, ValueError):
            return
        for raw_object in values:
            if len(seen) >= PDF_XOBJECT_MAX_OBJECTS:
                return
            visual = _resolve_pdf_object(raw_object)
            indirect_number = getattr(raw_object, "idnum", None)
            generation = getattr(raw_object, "generation", None)
            key = (
                ("indirect", indirect_number, generation)
                if indirect_number is not None
                else ("direct", id(visual))
            )
            if key in seen:
                continue
            seen.add(key)
            try:
                subtype = str(visual.get("/Subtype"))
            except (AttributeError, TypeError, ValueError):
                continue
            if subtype == "/Image":
                try:
                    width = int(visual.get("/Width") or 0)
                    height = int(visual.get("/Height") or 0)
                except (AttributeError, TypeError, ValueError):
                    continue
                if (
                    width > 0
                    and height > 0
                    and width * height >= PDF_LARGE_IMAGE_MIN_PIXELS
                ):
                    large_images += 1
                continue
            if subtype != "/Form":
                continue
            try:
                nested_resources = _resolve_pdf_object(
                    visual.get("/Resources")
                )
                nested_xobjects = nested_resources.get("/XObject")
            except (AttributeError, KeyError, TypeError, ValueError):
                continue
            if nested_xobjects is not None:
                visit(nested_xobjects, depth + 1)

    visit(xobjects, 0)
    return large_images


def _pdf_vector_operator_count(page: Any) -> int:
    try:
        content = page.get_contents()
        operations = list(getattr(content, "operations", []) or [])
    except (AttributeError, KeyError, TypeError, ValueError):
        return 0
    return sum(
        1
        for _, operator in operations
        if operator in _PDF_VECTOR_OPERATORS
    )


def _pdf_mixed_visual_reason(page: Any, text: str) -> dict[str, Any] | None:
    large_images = _pdf_large_image_count(page)
    vector_operators = _pdf_vector_operator_count(page)
    lowered = text.casefold()
    visual_unit_value_count = len(_PDF_VISUAL_UNIT_PATTERN.findall(text))
    keyword = next(
        (item for item in _PDF_VISUAL_KEYWORDS if item in lowered),
        None,
    )
    reason: str | None = None
    if large_images:
        reason = "large_embedded_image"
    elif (
        vector_operators >= PDF_VECTOR_VISUAL_OPERATOR_THRESHOLD
        and keyword is not None
    ):
        reason = "keyword_with_vector_drawing"
    elif (
        vector_operators >= PDF_VECTOR_VISUAL_HIGH_THRESHOLD
        and visual_unit_value_count >= 2
    ):
        reason = "unit_values_with_dense_vector_drawing"
    if reason is None:
        return None
    return {
        "reason": reason,
        "large_image_count": large_images,
        "vector_operator_count": vector_operators,
        "visual_unit_value_count": visual_unit_value_count,
        "keyword": keyword,
    }


def parse_pdf_document(path: Path, relative_path: str, file_hash: str) -> PdfParseResult:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("读取 PDF 需要安装 pypdf：pip install -e '.[documents]'") from exc

    reader = PdfReader(path)
    page_count = len(reader.pages)
    if page_count > MAX_PDF_PAGES:
        raise ValueError(f"PDF pages exceed safety limit: {page_count} > {MAX_PDF_PAGES}")
    blocks: list[SourceBlock] = []
    scanned_page_numbers: list[int] = []
    mixed_visual_page_numbers: list[int] = []
    visual_page_reasons: list[dict[str, Any]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            scanned_page_numbers.append(page_number)
            visual_page_reasons.append(
                {
                    "page": page_number,
                    "page_mode": "scanned",
                    "reason": "no_text_layer",
                }
            )
            continue
        for block_index, paragraph in enumerate(part.strip() for part in text.split("\n") if part.strip()):
            blocks.append(
                _make_block(
                    relative_path=relative_path,
                    source_kind="pdf_text",
                    file_hash=file_hash,
                    locator={"page": page_number, "text_block": block_index},
                    text=paragraph,
                )
            )
        mixed_reason = _pdf_mixed_visual_reason(page, text)
        if (
            mixed_reason is not None
            and len(mixed_visual_page_numbers) < MAX_PDF_MIXED_VISUAL_PAGES
        ):
            mixed_visual_page_numbers.append(page_number)
            visual_page_reasons.append(
                {
                    "page": page_number,
                    "page_mode": "mixed",
                    **mixed_reason,
                }
            )
    return PdfParseResult(
        blocks=tuple(blocks),
        scanned_page_numbers=tuple(scanned_page_numbers),
        page_count=page_count,
        mixed_visual_page_numbers=tuple(mixed_visual_page_numbers),
        visual_page_reasons=tuple(visual_page_reasons),
    )


def parse_pdf(path: Path, relative_path: str, file_hash: str) -> list[SourceBlock]:
    """Parse text-layer PDF pages while preserving the original public contract."""

    return list(parse_pdf_document(path, relative_path, file_hash).blocks)


def render_pdf_page(
    path: Path,
    page_number: int,
    output_dir: Path,
    *,
    scale: float = PDF_RENDER_SCALE,
) -> RenderedPdfPage:
    """Render one 1-based PDF page to a temporary PNG with bounded pixels.

    The caller owns ``output_dir`` and is responsible for removing it. The
    engine uses ``TemporaryDirectory`` so rendered customer pages are removed
    on success and on every exception path.
    """

    try:
        import pypdfium2 as pdfium
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError(
            "渲染扫描 PDF 需要安装 pypdfium2：pip install -e '.[documents]'"
        ) from exc

    if page_number < 1:
        raise ValueError("PDF page_number must be 1-based and positive.")
    if not 0.25 <= scale <= 4.0:
        raise ValueError("PDF render scale must be between 0.25 and 4.0.")

    output_dir.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(path))
    try:
        if page_number > len(document):
            raise ValueError(f"PDF page {page_number} does not exist.")
        page = document[page_number - 1]
        try:
            width_points, height_points = page.get_size()
            expected_width = max(1, ceil(float(width_points) * scale))
            expected_height = max(1, ceil(float(height_points) * scale))
            if expected_width * expected_height > MAX_PDF_RENDER_PIXELS:
                raise ValueError(
                    "Rendered PDF page would exceed the pixel safety limit: "
                    f"{expected_width}x{expected_height}"
                )
            bitmap = page.render(scale=scale)
            try:
                image = bitmap.to_pil()
                try:
                    pixel_width, pixel_height = image.size
                    image_path = output_dir / f"page-{page_number:04d}.png"
                    image.save(image_path, format="PNG")
                finally:
                    image.close()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()

    return RenderedPdfPage(
        page_number=page_number,
        image_path=image_path,
        pixel_width=int(pixel_width),
        pixel_height=int(pixel_height),
    )


def _bounded_pdfium_object_bounds(
    visual: Any,
    *,
    page_width: float,
    page_height: float,
) -> tuple[float, float, float, float] | None:
    try:
        raw_bounds = tuple(float(value) for value in visual.get_bounds())
    except (AttributeError, TypeError, ValueError):
        return None
    if len(raw_bounds) != 4 or any(
        not isfinite(value) for value in raw_bounds
    ):
        return None
    left, bottom, right, top = raw_bounds
    left = max(0.0, min(page_width, left))
    bottom = max(0.0, min(page_height, bottom))
    right = max(0.0, min(page_width, right))
    top = max(0.0, min(page_height, top))
    if left >= right or bottom >= top:
        return None
    return left, bottom, right, top


def _pdfium_large_image(visual: Any) -> bool:
    try:
        width, height = visual.get_px_size()
        width = int(width)
        height = int(height)
    except (AttributeError, TypeError, ValueError):
        return False
    return (
        width > 0
        and height > 0
        and width * height >= PDF_LARGE_IMAGE_MIN_PIXELS
    )


def _reliable_pdf_roi_placement(
    bounds: tuple[float, float, float, float],
    *,
    page_width: float,
    page_height: float,
) -> bool:
    left, bottom, right, top = bounds
    width = right - left
    height = top - bottom
    area_ratio = (width * height) / (page_width * page_height)
    if (
        area_ratio < PDF_VISUAL_ROI_MIN_PLACED_AREA_RATIO
        or width / page_width < PDF_VISUAL_ROI_MIN_PLACED_DIMENSION_RATIO
        or height / page_height < PDF_VISUAL_ROI_MIN_PLACED_DIMENSION_RATIO
    ):
        return False
    in_footer = top <= page_height * PDF_VISUAL_ROI_HEADER_FOOTER_RATIO
    in_header = (
        bottom
        >= page_height * (1.0 - PDF_VISUAL_ROI_HEADER_FOOTER_RATIO)
    )
    if (in_header or in_footer) and area_ratio < 0.05:
        return False
    return True


def _union_pdf_bounds(
    bounds: Iterable[tuple[float, float, float, float]],
) -> tuple[float, float, float, float]:
    rows = tuple(bounds)
    return (
        min(row[0] for row in rows),
        min(row[1] for row in rows),
        max(row[2] for row in rows),
        max(row[3] for row in rows),
    )


def _expanded_pdf_bounds(
    bounds: tuple[float, float, float, float],
    *,
    margin: float,
    page_width: float,
    page_height: float,
) -> tuple[float, float, float, float]:
    left, bottom, right, top = bounds
    return (
        max(0.0, left - margin),
        max(0.0, bottom - margin),
        min(page_width, right + margin),
        min(page_height, top + margin),
    )


def _pdf_bounds_intersect(
    first: tuple[float, float, float, float],
    second: tuple[float, float, float, float],
) -> bool:
    return not (
        first[2] < second[0]
        or first[0] > second[2]
        or first[3] < second[1]
        or first[1] > second[3]
    )


def _close_pdfium_objects(objects: Iterable[Any]) -> None:
    for visual in objects:
        try:
            visual.close()
        except (AttributeError, RuntimeError):
            pass


def _select_pdf_visual_region(
    page: Any,
    pdfium: Any,
    *,
    page_width: float,
    page_height: float,
) -> tuple[
    tuple[float, float, float, float] | None,
    float,
    str | None,
]:
    """Select a conservative raster ROI using page-space PDFium objects."""

    try:
        rotation = int(page.get_rotation())
    except (AttributeError, TypeError, ValueError):
        return None, 1.0, "page_rotation_unavailable"
    if rotation != 0:
        return None, 1.0, "page_rotation_not_supported"

    raw = getattr(pdfium, "raw", None)
    image_type = getattr(raw, "FPDF_PAGEOBJ_IMAGE", None)
    form_type = getattr(raw, "FPDF_PAGEOBJ_FORM", None)
    text_type = getattr(raw, "FPDF_PAGEOBJ_TEXT", None)
    if image_type is None or form_type is None or text_type is None:
        return None, 1.0, "pdfium_object_api_unavailable"

    candidates: list[tuple[float, float, float, float]] = []
    neighboring_visuals: list[
        tuple[float, float, float, float]
    ] = []
    opened_objects: list[Any] = []
    try:
        direct_images = list(
            islice(
                page.get_objects(filter=[image_type], max_depth=0),
                PDF_XOBJECT_MAX_OBJECTS,
            )
        )
        opened_objects.extend(direct_images)
        for image in direct_images:
            bounds = _bounded_pdfium_object_bounds(
                image,
                page_width=page_width,
                page_height=page_height,
            )
            if bounds is not None and _reliable_pdf_roi_placement(
                bounds,
                page_width=page_width,
                page_height=page_height,
            ):
                neighboring_visuals.append(bounds)
                if _pdfium_large_image(image):
                    candidates.append(bounds)

        forms = list(
            islice(
                page.get_objects(filter=[form_type], max_depth=0),
                PDF_XOBJECT_MAX_OBJECTS,
            )
        )
        opened_objects.extend(forms)
        descendant_budget = PDF_XOBJECT_MAX_OBJECTS
        for form_index, form in enumerate(forms):
            bounds = _bounded_pdfium_object_bounds(
                form,
                page_width=page_width,
                page_height=page_height,
            )
            if bounds is None or not _reliable_pdf_roi_placement(
                bounds,
                page_width=page_width,
                page_height=page_height,
            ):
                continue
            neighboring_visuals.append(bounds)
            if descendant_budget <= 0:
                descendant_images = []
            else:
                remaining_forms = max(1, len(forms) - form_index)
                per_form_budget = max(
                    1,
                    descendant_budget // remaining_forms,
                )
                descendant_images = list(
                    islice(
                        page.get_objects(
                            filter=[image_type],
                            form=form,
                            max_depth=PDF_XOBJECT_MAX_DEPTH,
                        ),
                        per_form_budget,
                    )
                )
            opened_objects.extend(descendant_images)
            descendant_budget = max(
                0,
                descendant_budget - len(descendant_images),
            )
            if any(
                _pdfium_large_image(image)
                for image in descendant_images
            ):
                candidates.append(bounds)

        if not candidates:
            return None, 1.0, "no_reliable_raster_region"

        text_objects = list(
            islice(
                page.get_objects(filter=[text_type], max_depth=0),
                PDF_XOBJECT_MAX_OBJECTS,
            )
        )
        opened_objects.extend(text_objects)
        text_bounds = [
            bounds
            for text in text_objects
            if (
                bounds := _bounded_pdfium_object_bounds(
                    text,
                    page_width=page_width,
                    page_height=page_height,
                )
            )
            is not None
        ]

        content_region = _union_pdf_bounds(candidates)
        # Absorb nearby text and any visual object that would otherwise be
        # sliced by the safety margin. This avoids returning a crop containing
        # half an adjacent chart or diagram.
        for _ in range(3):
            snap_region = _expanded_pdf_bounds(
                content_region,
                margin=PDF_VISUAL_ROI_TEXT_SNAP_POINTS,
                page_width=page_width,
                page_height=page_height,
            )
            snapped_text = [
                bounds
                for bounds in text_bounds
                if _pdf_bounds_intersect(bounds, snap_region)
            ]
            if snapped_text:
                content_region = _union_pdf_bounds(
                    (content_region, *snapped_text)
                )
            margin_region = _expanded_pdf_bounds(
                content_region,
                margin=PDF_VISUAL_ROI_MARGIN_POINTS,
                page_width=page_width,
                page_height=page_height,
            )
            touching_visuals = [
                bounds
                for bounds in neighboring_visuals
                if _pdf_bounds_intersect(bounds, margin_region)
            ]
            expanded_content = (
                _union_pdf_bounds(
                    (content_region, *touching_visuals)
                )
                if touching_visuals
                else content_region
            )
            if expanded_content == content_region:
                break
            content_region = expanded_content
        region = _expanded_pdf_bounds(
            content_region,
            margin=PDF_VISUAL_ROI_MARGIN_POINTS,
            page_width=page_width,
            page_height=page_height,
        )
    finally:
        _close_pdfium_objects(reversed(opened_objects))

    left, bottom, right, top = region
    crop_ratio = (
        (right - left) * (top - bottom)
        / (page_width * page_height)
    )
    if crop_ratio >= PDF_VISUAL_ROI_MAX_AREA_RATIO:
        return None, crop_ratio, "crop_area_ratio_limit"
    if 1.0 - crop_ratio < PDF_VISUAL_ROI_MIN_SAVINGS_RATIO:
        return None, crop_ratio, "crop_savings_below_threshold"
    return region, crop_ratio, None


def render_pdf_visual_region(
    path: Path,
    page_number: int,
    output_dir: Path,
    *,
    scale: float = PDF_RENDER_SCALE,
) -> RenderedPdfVisualRegion:
    """Render a conservative raster region, failing open to the full page.

    ROI v1 intentionally handles only unrotated, reliable top-level images or
    top-level forms that contain a large raster image. Vector-only drawings,
    scanned pages, uncertain object trees, and low-savings crops remain full
    page. OCR edge guarding is deliberately left to the reader integration.
    """

    if page_number < 1:
        raise ValueError("PDF page_number must be 1-based and positive.")
    if not 0.25 <= scale <= 4.0:
        raise ValueError("PDF render scale must be between 0.25 and 4.0.")
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError(
            "渲染 PDF 视觉区域需要安装 pypdfium2："
            "pip install -e '.[documents]'"
        ) from exc

    output_dir.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(path))
    fallback_reason: str | None = None
    page_width = 0.0
    page_height = 0.0
    crop_ratio = 1.0
    try:
        if page_number > len(document):
            raise ValueError(f"PDF page {page_number} does not exist.")
        page = document[page_number - 1]
        try:
            raw_width, raw_height = page.get_size()
            page_width = float(raw_width)
            page_height = float(raw_height)
            if (
                not isfinite(page_width)
                or not isfinite(page_height)
                or page_width <= 0
                or page_height <= 0
            ):
                fallback_reason = "invalid_page_size"
                region = None
            else:
                try:
                    raw_crop_box = tuple(
                        float(value) for value in page.get_cropbox()
                    )
                    if (
                        len(raw_crop_box) != 4
                        or not all(isfinite(value) for value in raw_crop_box)
                    ):
                        raise ValueError("invalid crop box")
                except (AttributeError, RuntimeError, TypeError, ValueError):
                    region = None
                    fallback_reason = "page_cropbox_unavailable"
                else:
                    crop_left, crop_bottom, crop_right, crop_top = (
                        raw_crop_box
                    )
                    tolerance = max(
                        0.001,
                        page_width * 0.000001,
                        page_height * 0.000001,
                    )
                    if (
                        abs(crop_left) > tolerance
                        or abs(crop_bottom) > tolerance
                        or abs(crop_right - page_width) > tolerance
                        or abs(crop_top - page_height) > tolerance
                    ):
                        # PDFium object bounds may remain in media-page
                        # coordinates while get_size/render crop coordinates
                        # are CropBox-relative. ROI v1 therefore fails open
                        # instead of risking a silently shifted crop.
                        region = None
                        fallback_reason = (
                            "nonzero_or_offset_page_cropbox_not_supported"
                        )
                    else:
                        try:
                            (
                                region,
                                crop_ratio,
                                fallback_reason,
                            ) = _select_pdf_visual_region(
                                page,
                                pdfium,
                                page_width=page_width,
                                page_height=page_height,
                            )
                        except Exception as exc:
                            region = None
                            fallback_reason = (
                                "roi_detection_failed:"
                                f"{type(exc).__name__}"
                            )

            if region is not None:
                left, bottom, right, top = region
                expected_width = max(1, ceil((right - left) * scale))
                expected_height = max(1, ceil((top - bottom) * scale))
                if (
                    expected_width * expected_height
                    > MAX_PDF_RENDER_PIXELS
                ):
                    fallback_reason = "roi_pixel_limit"
                else:
                    image_path = (
                        output_dir / f"page-{page_number:04d}-roi.png"
                    )
                    try:
                        bitmap = page.render(
                            scale=scale,
                            crop=(
                                left,
                                bottom,
                                page_width - right,
                                page_height - top,
                            ),
                        )
                        try:
                            image = bitmap.to_pil()
                            try:
                                pixel_width, pixel_height = image.size
                                image.save(image_path, format="PNG")
                            finally:
                                image.close()
                        finally:
                            bitmap.close()
                    except Exception as exc:
                        image_path.unlink(missing_ok=True)
                        fallback_reason = (
                            f"roi_render_failed:{type(exc).__name__}"
                        )
                    else:
                        return RenderedPdfVisualRegion(
                            page_number=page_number,
                            image_path=image_path,
                            pixel_width=int(pixel_width),
                            pixel_height=int(pixel_height),
                            crop_box_pdf=tuple(
                                round(value, 6) for value in region
                            ),
                            crop_ratio=round(crop_ratio, 6),
                            strategy="pdfium_top_level_raster_union_v1",
                            fallback_reason=None,
                            page_width_pdf=round(page_width, 6),
                            page_height_pdf=round(page_height, 6),
                        )
        finally:
            page.close()
    finally:
        document.close()

    rendered = render_pdf_page(
        path,
        page_number,
        output_dir,
        scale=scale,
    )
    full_width = (
        page_width
        if page_width > 0
        else rendered.pixel_width / scale
    )
    full_height = (
        page_height
        if page_height > 0
        else rendered.pixel_height / scale
    )
    return RenderedPdfVisualRegion(
        page_number=rendered.page_number,
        image_path=rendered.image_path,
        pixel_width=rendered.pixel_width,
        pixel_height=rendered.pixel_height,
        crop_box_pdf=(0.0, 0.0, full_width, full_height),
        crop_ratio=1.0,
        strategy="full_page",
        fallback_reason=(
            fallback_reason or "no_reliable_raster_region"
        ),
        page_width_pdf=round(full_width, 6),
        page_height_pdf=round(full_height, 6),
    )


def parse_file(path: Path, root: Path) -> list[SourceBlock]:
    relative_path = path.relative_to(root).as_posix()
    file_hash = sha256_file(path)
    lowered = path.name.casefold()
    if lowered.endswith(".ocr.json"):
        return parse_ocr_sidecar(path, relative_path, file_hash)
    suffix = path.suffix.casefold()
    if suffix in TEXT_EXTENSIONS:
        return parse_text(path, relative_path, file_hash)
    if suffix == ".csv":
        return parse_csv(path, relative_path, file_hash)
    if suffix == ".json":
        return parse_json(path, relative_path, file_hash)
    if suffix == ".docx":
        return parse_docx(path, relative_path, file_hash)
    if suffix == ".xlsx":
        return parse_xlsx(path, relative_path, file_hash)
    if suffix == ".pdf":
        return parse_pdf(path, relative_path, file_hash)
    if suffix in IMAGE_EXTENSIONS:
        return []
    raise ValueError(f"Unsupported file type: {relative_path}")


def discover_files(root: Path) -> list[Path]:
    canonical_root = root.expanduser().resolve(strict=True)
    files: set[Path] = set()
    for path in canonical_root.rglob("*"):
        lexical_relative = path.relative_to(canonical_root)
        if any(part.startswith(".") for part in lexical_relative.parts):
            continue
        try:
            resolved = path.resolve(strict=True)
            resolved_relative = resolved.relative_to(canonical_root)
        except (OSError, RuntimeError, ValueError) as exc:
            raise ValueError(
                f"Input path escapes the selected directory or is unsafe: {lexical_relative.as_posix()}"
            ) from exc
        if not resolved.is_file():
            continue
        if any(part.startswith(".") for part in resolved_relative.parts):
            continue
        name = resolved.name.casefold()
        if name.endswith(".ocr.json") or resolved.suffix.casefold() in SUPPORTED_EXTENSIONS:
            files.add(resolved)
    return sorted(files)
