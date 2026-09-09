"""Review presentation and portable, local Excel delivery; never accepts decisions."""
from __future__ import annotations

from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import tempfile

from . import workflow
from .confirmation import ConfirmationRequestError
from .graph import build_graph
from .models import FactCandidate


SCOPES = {"input": "输入", "output": "输出", "rated": "额定", "nominal": "标称",
          "min": "最小", "max": "最大", "typical": "典型", "net": "净重", "gross": "毛重",
          "minimum": "最小", "maximum": "最大", "unspecified": "未限定",
          "operating": "工作/运行", "storage": "储存"}
STATUSES = {"pending": "未作人工决定", "confirmed": "已人工确认", "rejected": "已拒绝", "stale": "确认已失效",
            "source_changed": "来源已变化"}
CLASSIFICATIONS = {"strong_conflict": "数值冲突", "converted_match": "换算一致", "exact_match": "表达一致",
    "insufficient_evidence": "证据不足", "compatible_expression": "表达待统一", "likely_version_update": "可能版本不同",
    "semantic_scope_split": "不同口径", "identity_ambiguous": "先确认商品归属",
    "source_changed": "先更新资料", "decided": "已作人工决定"}


def scope_label(scope):
    return " / ".join(SCOPES.get(part.removeprefix("rating:"), part.replace("profile:", "模式：").replace("variant:", "变体："))
                      for part in (scope or "unspecified").split("|"))


def _unit(unit):
    return "件" if unit == "count" else unit or ""


def group_candidates(rows, hashes):
    """Use the same facts as analysis; source validation never chooses a value."""
    candidates = [FactCandidate.from_dict(row) for row in rows]
    for candidate in candidates:
        candidate.source_current = (candidate.source_current and
                                    hashes.get(candidate.source_file) == candidate.file_hash)
    groups = []
    _, fact_groups, _ = build_graph(candidates)
    for graph in fact_groups:
        groups.append({
            "group_id": graph.group_id, "product_id": graph.product_id,
            "product_label": graph.product_label, "field": graph.field,
            "label": graph.field_label, "scope": graph.scope, "scope_label": scope_label(graph.scope),
            "classification": "source_changed" if not graph.current else graph.classification,
            "severity": {"conflict": "block", "pending_confirmation": "review", "verified": "pass"}[graph.fact_status],
            "reason": graph.reason, "source_count": graph.independent_source_count,
            "candidate_ids": graph.candidate_ids, "fact_status": graph.fact_status,
            "verification_method": graph.verification_method, "selected_value": graph.selected_value,
            "selected_unit": graph.selected_unit, "review_reason_code": graph.review_reason_code,
            "human_approved": graph.human_approved, "current": graph.current,
        })
    return sorted(groups, key=lambda g: ({"conflict": 0, "pending_confirmation": 1, "verified": 2}[g["fact_status"]],
                                        g.get("product_id") or "", g["field"], g["scope"] or ""))

def snapshot(output, product):
    from .conversation import _snapshot
    rows, hashes, _ = _snapshot(output, product, sorted({c["field"] for c in product["candidates"]}))
    version = workflow.digest({"session": product["run_summary"]["session_id"], "rows": rows, "hashes": hashes})
    return rows, hashes, version


def _value(value, field=None):
    if isinstance(value, list):
        return (" × " if field == "dimensions" else " 至 ").join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _position(locator):
    labels = {"page": "页", "line": "行", "paragraph": "段落", "sheet": "工作表", "cell": "单元格",
              "row": "行", "column": "列", "table": "表格"}
    parts = [f"{label} {locator[key]}" for key, label in labels.items() if locator.get(key) is not None]
    if locator.get("bbox_1000"):
        parts.append("图片近似区域 " + str(locator["bbox_1000"]))
    elif "image" in locator:
        parts.append("图片文字（无精确坐标）")
    return "；".join(parts) or "见来源原文"


def export_review(output_dir, *, session_id=None):
    """Export a fact-first local workbook; keep evidence in the final sheet."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise ConfirmationRequestError("Excel 导出依赖未安装，请运行 scripts\\install-env.ps1。") from exc
    output, product, _ = workflow.context(output_dir, session_id)
    rows, hashes, version = snapshot(output, product)
    if not rows:
        raise ConfirmationRequestError("当前没有参数候选，请先分析包含商品参数的资料。")
    groups = group_candidates(rows, hashes)
    by_id = {c["candidate_id"]: c for c in rows}
    generated = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %z")
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, subtitle, headings, records, widths):
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.append([title]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headings))
        ws.cell(1, 1).font = Font(name="微软雅黑", size=16, bold=True, color="183D4D")
        ws.row_dimensions[1].height = 35
        ws.append([subtitle]); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headings))
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(2, 1).font = Font(name="微软雅黑", size=10, color="506675")
        ws.row_dimensions[2].height = 42
        ws.append(["生成时间：" + generated]); ws.row_dimensions[3].height = 23
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headings))
        ws.append(headings)
        for record in records:
            ws.append(record)
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.font = Font(name="微软雅黑", size=10, color="233B49")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, str):
                    # Data from documents must remain literal, including =/+/@.
                    cell.data_type = "s"
            if row[0].row >= 5:
                ws.row_dimensions[row[0].row].height = min(240, max(42, 16 * max(str(c.value or "").count("\n") + 1 for c in row)))
        for cell in ws[4]:
            cell.fill = PatternFill("solid", fgColor="183D4D")
            cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True, size=10)
        ws.row_dimensions[4].height = 28
        if records:
            table = Table(displayName=f"ReviewTable{len(wb.worksheets)}", ref=f"A4:{ws.cell(ws.max_row, len(headings)).coordinate}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        ws.freeze_panes = "C5"
        from openpyxl.utils import get_column_letter
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:4"
        ws.print_options.horizontalCentered = True
        return ws

    overview, conflicts, pending, selected, details = [], [], [], [], []
    human_count = 0
    status_labels = {"conflict": "冲突", "pending_confirmation": "待确认", "verified": "已确认"}
    methods = {"single_value": "参数明确，无冲突", "cross_source_exact": "多来源一致",
               "cross_source_converted": "单位换算后一致", "human_confirmed": "人工确认",
               "human_resolved_conflict": "人工选择采用值"}
    for group in groups:
        items = [by_id[cid] for cid in group["candidate_ids"]]
        live = [c for c in items if c["status"] != "rejected" and hashes.get(c["source_file"]) == c["file_hash"]]
        values = {}
        for c in live:
            key = workflow.digest([c["normalized_value"], c.get("normalized_unit")])
            values.setdefault(key, {"value": _value(c["normalized_value"], c["field"]), "unit": _unit(c.get("normalized_unit")), "sources": set()})
            values[key]["sources"].add(c["source_file"])
        alternatives = "\n".join(f'{v["value"]} {v["unit"]}'.strip() for v in values.values())
        source_values = "\n".join("{} {}（{}）".format(v["value"], v["unit"], "、".join(sorted(v["sources"]))) for v in values.values())
        source_summary = "、".join(sorted({c["source_file"] for c in live}))
        context = [group.get("product_label") or "未归属商品", group["label"], group["scope_label"]]
        status = group["fact_status"]
        shown_value = _value(group["selected_value"], group["field"]) if status == "verified" else alternatives
        unit = _unit(group["selected_unit"]) if status == "verified" else ""
        reason = "" if status == "verified" else group["reason"]
        overview.append([*context, status_labels[status], shown_value, unit, reason])
        if status == "conflict":
            conflicts.append([*context, source_values, group["source_count"], "请选择采用值，或补充说明。"])
        elif status == "pending_confirmation":
            pending.append([*context, alternatives, group["reason"], group["source_count"], source_summary])
        else:
            human_count += int(group["human_approved"])
            selected.append([*context, shown_value, unit, methods.get(group["verification_method"], ""),
                             group["source_count"], source_summary, "是" if group["human_approved"] else "否", "是"])
        for c in items:
            current = hashes.get(c["source_file"]) == c["file_hash"] and c.get("source_current", True)
            decision = c["status"] if current else "source_changed"
            details.append([group["group_id"], c["candidate_id"], *context, _value(c["normalized_value"], c["field"]),
                            _unit(c.get("normalized_unit")), c["raw_value"], c["source_file"],
                            _position(c.get("locator", {})), c["raw_text"], c["extraction_method"],
                            min(c["recognition_confidence"], c["mapping_confidence"]), STATUSES.get(decision, decision)])

    main = sheet("核验总览", "清楚且无冲突的参数直接列入已确认。只需处理冲突和待确认项目。",
                 ["商品", "参数", "口径", "状态", "当前值", "单位", "需要处理"], overview,
                 [28, 18, 20, 14, 28, 12, 58])
    for index, group in enumerate(groups, 5):
        color = "E8EAED" if not group["current"] or group["review_reason_code"] == "all_rejected" else {
            "conflict": "FBE4DE", "pending_confirmation": "FFF3D6", "verified": "E5F1EA"}[group["fact_status"]]
        main.cell(index, 4).fill = PatternFill("solid", fgColor=color)
    sheet("冲突", "不同值并列展示，不按来源多数自动选择。",
          ["商品", "参数", "口径", "冲突值及来源", "独立来源数", "建议操作"],
          conflicts, [28, 18, 20, 65, 16, 45])
    sheet("待确认事实", "只列出归属、口径、识别或来源变化等实际疑问。",
          ["商品", "参数", "口径", "待确认值", "待确认原因", "独立来源数", "来源摘要"],
          pending, [28, 18, 20, 28, 55, 16, 48])
    sheet("已确认事实", "包含自动核验和人工确认。自动核验不代表用户已正式批准。",
          ["商品", "参数", "口径", "已确认值", "单位", "确认方式", "独立来源数", "来源摘要", "是否人工确认", "当前有效"],
          selected, [28, 18, 20, 28, 12, 24, 16, 48, 18, 14])
    evidence = sheet("证据明细", "仅供追溯。修改 Excel 不会改变系统的确认记录。",
          ["事实组 ID", "候选 ID", "商品", "参数", "口径", "标准值", "单位", "原始值", "来源文件", "位置",
           "证据原文", "提取方式", "置信度", "人工决策"],
          details, [30, 30, 28, 18, 20, 24, 12, 28, 40, 32, 65, 30, 14, 20])
    for row in evidence.iter_rows(min_row=5):
        row[12].number_format = "0%"
        if row[13].value in {"已拒绝", "确认已失效", "来源已变化"}:
            row[13].fill = PatternFill("solid", fgColor="E8EAED")

    path = output / "product-review.xlsx"
    if path.exists():
        workflow.safe_file(path, max_bytes=20_000_000)
    descriptor, temporary_name = tempfile.mkstemp(dir=output, prefix=".product-review-", suffix=".xlsx")
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        wb.save(temporary)
        _, current_product, _ = workflow.context(output, product["run_summary"]["session_id"])
        if snapshot(output, current_product)[2] != version:
            raise ConfirmationRequestError("资料或确认结果在导出期间改变，请重新导出。")
        os.replace(temporary, path)
    except PermissionError as exc:
        raise ConfirmationRequestError("核验表正在被 Excel 占用，请关闭该文件后重新导出。") from exc
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)
    session = product["run_summary"]["session_id"]
    manifest = workflow._manifest(output, session)
    artifact_id = workflow.digest(str(path))[:24]
    manifest["deliverables"][artifact_id] = {"artifact_id": artifact_id, "kind": "review_workbook",
        "name": "商品核验表", "path": str(path), "file_hash": hashlib.sha256(path.read_bytes()).hexdigest(),
        "bundle_id": "local-review", "bundle_version": version, "review_version": version,
        "checked_at": workflow.now(), "revision": manifest["deliverables"].get(artifact_id, {}).get("revision", 0) + 1,
        "references": [], "check": {"status": "review_snapshot", "finding_count": sum(g["severity"] == "block" for g in groups)}}
    from .state import atomic_write_json
    atomic_write_json(output / workflow.WORKFLOW_FILE, manifest)
    return {"path": str(path), "artifact_id": artifact_id, "status": "review_workbook_ready",
        "candidate_count": len(rows), "confirmed_row_count": human_count, "verified_row_count": len(selected), "fact_count": len(groups), "needs_reanalysis": any(hashes.get(c["source_file"]) != c["file_hash"] for c in rows),
        "processing": "local", "network_sent": False, "model_called": False,
        "next_action": "open_local_review_workbook", "contains_source_evidence": True}
